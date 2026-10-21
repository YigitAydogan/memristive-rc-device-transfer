"""
Memristive reservoir portability, v7: a measured device pair.

The TiOx dataset (Wang et al., J. Phys. D: Appl. Phys. 59 015101, 2026; dataset 10.5258/SOTON/D3702) contains two
pulse-train runs with the same nominal protocol (4 V, 100 us pulses, 8 us apart): run A (Fig 2c) and run B (Fig 2d).
Their currents differ 4.5-fold at rest and their responses grow differently. Using the v6 calibration, run A and run B
become two measurement-calibrated devices. This script tests readout transfer between them, in both directions, with
the same drive protocol and with each device's own interface, at nominal and at measured noise, and places the pair on
a simulated spread axis (5% to 150%) by feature correlation.

Requires the finished v6 folder (calibration) and v5 folder (universal operating point); the measured files must still
be in MyDrive/memristive_device_generalization/data/measured. Paste into one Colab cell with Drive mounted. The script
does not mount Drive. Units are checkpointed.
"""

from __future__ import annotations

import hashlib
import importlib
import json
import math
import os
import platform
import subprocess
import sys
import time
import traceback
import warnings
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import FixedLocator, FuncFormatter
from mpl_toolkits.axes_grid1 import make_axes_locatable
from scipy.stats import spearmanr, wilcoxon
from sklearn.model_selection import StratifiedKFold
from sklearn.neighbors import KNeighborsClassifier

warnings.filterwarnings("ignore", category=RuntimeWarning)
warnings.filterwarnings("ignore", category=UserWarning)
warnings.filterwarnings("ignore", category=FutureWarning)

# =============================================================================
# Configuration
# =============================================================================

SCRIPT_VERSION = "7.0.0"
SMOKE = os.environ.get("PORTV2_SMOKE", "0").strip() == "1"
IDEAL = os.environ.get("PORTV2_IDEAL", "0").strip() == "1"
RUN_MINIROCKET = os.environ.get("PORTV2_MINIROCKET", "1").strip() == "1" and not SMOKE
DATASET_FILTER = [d.strip() for d in os.environ.get("PORTV2_DATASETS", "").split(",") if d.strip()]

DEFAULT_DRIVE_ROOT = Path("/content/drive/MyDrive/memristive_device_generalization")
LOCAL_ROOT = Path("/content/memristive_device_generalization")
ENV_ROOT = os.environ.get("MEMRISTIVE_PROJECT_ROOT", "").strip()

# Datasets previously examined in the earlier project (reported as their own tier)
# and datasets never used before (fresh tier).
PRIOR_DATASETS = [
    "ECG5000", "FordA", "Wafer", "UWaveGestureLibraryX",
    "CricketX", "RefrigerationDevices", "Earthquakes", "StarLightCurves",
]
FRESH_DATASETS = [
    "GunPoint", "TwoPatterns", "SyntheticControl", "Trace", "FaceAll",
    "Strawberry", "Yoga", "ECG200", "Plane", "Symbols",
]

SEEDS = [11, 22, 33] if not SMOKE else [11]
GAINS = np.array([0.20, 0.35, 0.50, 0.65, 0.80, 1.00]) if not SMOKE else np.array([0.35, 1.00])
FRACS = np.array([0.03, 0.06, 0.125, 0.25, 0.50, 1.00]) if not SMOKE else np.array([0.06, 0.50])
N_BINS = 24
ALPHA_REL = np.array([1e-4, 1e-3, 1e-2, 1e-1, 1.0])
K_SHOTS = [1, 2, 5, 10] if not SMOKE else [1, 2]
LAMBDA_PRIOR = 10.0          # warm-start and scratch k-shot shrinkage, in sample units
CORAL_EPS = 0.05
FIR_TAPS = 24
QUAD_TAPS = 12
PROBE_EPISODES = 50 if not SMOKE else 12
PROBE_LENGTH = 256 if not SMOKE else 128
PROBE_TRAIN_EPISODES = int(round(0.75 * PROBE_EPISODES))
CAP_STATE_LAGS = 24
CAP_MAX_DELAY_D1 = 39
CAP_MAX_DELAY_D2 = 14
CAP_MAX_DELAY_D3 = 9
CAP_THRESHOLD = 0.01
RETENTION_PER_SYMBOL = 0.90
SIGMA_D2D = 0.0 if IDEAL else 0.05
SIGMA_C2C = 0.0 if IDEAL else 0.002
SIGMA_READ = 0.0 if IDEAL else 0.01   # decades of read conductance
N_BOOT = 10000 if not SMOKE else 300
MASTER_SEED = 20261020
EPS = 1e-12

PRETTY_METHOD = {
    "zero_shot": "Zero shot",
    "moment": "Moment matching",
    "coral": "CORAL",
    "replay_fir": "Replay FIR",
    "probe_fir": "Probe FIR",
    "probe_quad": "Probe quadratic",
    "probe_fir_mm": "Probe FIR + moment",
    "full_refit": "Full refit",
}
LABEL_FREE = ["coral", "replay_fir", "probe_fir", "probe_quad", "probe_fir_mm"]

# =============================================================================
# Paths and logging
# =============================================================================


def choose_root() -> Path:
    if ENV_ROOT:
        return Path(ENV_ROOT).expanduser().resolve()
    if DEFAULT_DRIVE_ROOT.parent.exists():
        return DEFAULT_DRIVE_ROOT
    if "google.colab" in sys.modules:
        raise RuntimeError("Google Drive is not mounted at /content/drive/MyDrive. Mount it before running.")
    return LOCAL_ROOT


PROJECT_ROOT = choose_root()
RUN_NAME = "portability_v7_measured_pair" + ("_ideal" if IDEAL else "") + ("_smoke" if SMOKE else "")
ROOT = PROJECT_ROOT / RUN_NAME
DIRS = {k: ROOT / k for k in ["config", "units", "probe", "tables", "figures", "logs", "data"]}
for _d in DIRS.values():
    _d.mkdir(parents=True, exist_ok=True)
DATA_DIR = PROJECT_ROOT / "data" / "tsml_ucr"
DATA_DIR.mkdir(parents=True, exist_ok=True)
LOG_FILE = DIRS["logs"] / "run.log"


def log(msg: str) -> None:
    line = f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | {msg}"
    print(line, flush=True)
    with LOG_FILE.open("a", encoding="utf-8") as f:
        f.write(line + "\n")


def stable_seed(*parts: Any) -> int:
    h = hashlib.sha256("|".join(map(str, parts)).encode("utf-8")).digest()
    return int.from_bytes(h[:8], "little") % (2**63 - 1)


def save_json(obj: Any, path: Path) -> None:
    path.write_text(json.dumps(obj, indent=2, sort_keys=True, default=str), encoding="utf-8")


# =============================================================================
# Device families
# =============================================================================


@dataclass(frozen=True)
class Family:
    code: str
    name: str
    citation: str
    voltage_mode: str      # "bipolar" or "unipolar"
    v_pos: float
    v_mid: float
    discrete: bool
    volatile: bool
    dt_hint: float
    x_lo: float
    x_hi: float
    params: Dict[str, float]


FAMILIES: Dict[str, Family] = {
    "LD": Family("LD", "Linear ion drift", "Strukov et al., Nature 2008", "bipolar", 0.5, 0.0, False, False,
                 1e-3, 0.0, 1.0, {"R_on": 100.0, "R_off": 1e5, "D": 1e-8, "mu_v": 1e-14, "x0": 0.5}),
    "VT": Family("VT", "VTEAM (fitted SDC)", "Bednarz and Garda, Energies 2021", "bipolar", 0.5, 0.0, False, False,
                 1e-4, 0.0, 1.0, {"R_on": 1357.8, "R_off": 352196.0, "k_on": -997.642, "k_off": 0.2969,
                                  "alpha_on": 2.0, "alpha_off": 1.0, "v_on": -0.340661, "v_off": 0.044669,
                                  "x0": 0.00099943}),
    "VV": Family("VV", "V-VTEAM volatile", "Patni et al., IFETC 2024", "unipolar", 3.0, 1.6583, False, True,
                 1e-5, 0.0, 1.0, {"v_th": 1.744, "v_hold": 1.5726, "R_on": 3e4, "R_off": 15e9, "k": 650.0,
                                  "alpha": 0.09999, "tau": 1.0444, "beta": 2.14262, "x0": 0.0}),
    "ZH": Family("ZH", "Dynamic memristor (Zhong)", "Zhong et al., Nat. Commun. 2021", "bipolar", 3.0, 0.0, True,
                 True, 1.0, 0.0, 1.0, {"r": 0.99, "G0": 0.5, "Kp": 9.13, "Kn": 0.32, "alpha": 0.23, "x0": 0.5}),
    "DF": Family("DF", "Diffusive compact", "Cisternas Ferri et al., Materials 2019", "bipolar", 2.5, 0.0, False,
                 False, 1e-6, 0.0, 1.0, {"alpha_plus": 15.0, "alpha_minus": 15.0, "delta_plus": 0.2,
                                         "delta_minus": 0.2, "R_on": 35.0, "R_off": 9500.0, "v0": 0.3,
                                         "tau0": 0.01, "x0": 0.5, "lambda0": 0.5}),
    "DN": Family("DN", "Volatile dynamic (Daniels)", "Daniels et al., Adv. Electron. Mater. 2026", "unipolar", 0.8,
                 0.5, True, True, 1.0, 0.1, 1.0, {"alpha_i": 1e-8, "beta_i": 0.5, "gamma_i": 1e-5, "delta_i": 4.0,
                                                  "lambda": 1e3, "eta": 8.0, "x_max": 1.0, "x_min": 0.1,
                                                  "V_th": 0.6, "t_pulse": 1e-9, "tau": 10e-9, "x0": 0.1}),
}
FAMILY_CODES = list(FAMILIES.keys())
READ_VOLTAGE = 0.1
TX_LIKE = ("TX", "TB", "TS")   # measurement-calibrated multi-timescale volatile devices


@dataclass
class Instance:
    family: str
    role: str
    seed: int
    params: Dict[str, float]

    @property
    def iid(self) -> str:
        return f"{self.family}_{self.role}_s{self.seed}"


def _lognorm(v: float, rng: np.random.Generator) -> float:
    if SIGMA_D2D <= 0 or v == 0:
        return float(v)
    return float(math.copysign(abs(v) * math.exp(SIGMA_D2D * rng.standard_normal()), v))


def make_instance(fam: str, seed: int, role: str) -> Instance:
    p = dict(FAMILIES[fam].params)
    rng = np.random.default_rng(stable_seed("instance", fam, seed, role, MASTER_SEED))
    if fam == "LD":
        for k in ["R_on", "R_off", "mu_v"]:
            p[k] = _lognorm(p[k], rng)
    elif fam == "VT":
        for k in ["R_on", "R_off", "k_on", "k_off", "v_on", "v_off"]:
            p[k] = _lognorm(p[k], rng)
    elif fam == "VV":
        for k in ["R_on", "R_off", "k", "alpha", "tau", "beta", "v_th", "v_hold"]:
            p[k] = _lognorm(p[k], rng)
        p["v_hold"] = min(p["v_hold"], 0.97 * p["v_th"])
    elif fam == "ZH":
        leak = _lognorm(1.0 - p["r"], rng)
        p["r"] = float(np.clip(1.0 - leak, 0.8, 0.9999))
        for k in ["G0", "Kp", "Kn", "alpha"]:
            p[k] = _lognorm(p[k], rng)
        p["G0"] = float(np.clip(p["G0"], 0.05, 0.95))
    elif fam == "DF":
        for k in ["alpha_plus", "alpha_minus", "delta_plus", "delta_minus", "R_on", "R_off", "v0", "tau0"]:
            p[k] = _lognorm(p[k], rng)
    elif fam == "DN":
        for k in ["alpha_i", "beta_i", "gamma_i", "delta_i", "lambda", "eta", "V_th", "tau"]:
            p[k] = _lognorm(p[k], rng)
        p["V_th"] = float(np.clip(p["V_th"], 0.45, 0.75))
    elif fam in TX_LIKE:
        for k in ["m0", "m1", "m2", "k0", "k1", "k2", "tau0", "tau1", "tau2", "v_th", "alpha"]:
            p[k] = _lognorm(p[k], rng)
        p["v_th"] = float(min(p["v_th"], 0.95 * FAMILIES[fam].v_pos))
    return Instance(fam, role, seed, p)


def nominal_instance(fam: str) -> Instance:
    return Instance(fam, "nominal", -1, dict(FAMILIES[fam].params))


def u_to_voltage(fam: str, u: np.ndarray) -> np.ndarray:
    f = FAMILIES[fam]
    u = np.clip(u, -1.0, 1.0)
    if f.voltage_mode == "bipolar":
        return u * f.v_pos
    return np.where(u <= 0.0, (u + 1.0) * f.v_mid, f.v_mid + u * (f.v_pos - f.v_mid))


def rest_input(fam: str) -> float:
    return -1.0 if FAMILIES[fam].voltage_mode == "unipolar" else 0.0


def init_state(inst: Instance, shape: Tuple[int, ...]) -> Tuple[np.ndarray, Dict[str, np.ndarray]]:
    x = np.full(shape, inst.params["x0"], dtype=np.float64)
    aux: Dict[str, np.ndarray] = {}
    if inst.family == "VV":
        aux["age"] = np.zeros(shape, dtype=np.float64)
    if inst.family == "DF":
        aux["lam"] = np.full(shape, inst.params["lambda0"], dtype=np.float64)
    if inst.family in TX_LIKE:
        for j in range(3):
            aux[f"y{j}"] = np.zeros(shape, dtype=np.float64)
    return x, aux


def _vt_solve(target: np.ndarray, z0: np.ndarray) -> np.ndarray:
    """Invert G(z) = 0.5 (atanh z + atan z) via A = atanh z, where the map is well conditioned."""
    a = np.arctanh(np.clip(z0, 0.0, 1.0 - 1e-12))
    for _ in range(8):
        th = np.tanh(a)
        h = 0.5 * a + 0.5 * np.arctan(th) - target
        dh = 0.5 + 0.5 * (1.0 - th * th) / (1.0 + th * th)
        a = np.maximum(a - h / dh, 0.0)
    return np.clip(np.tanh(a), 0.0, 1.0 - 1e-12)


def advance(inst: Instance, x: np.ndarray, aux: Dict[str, np.ndarray], v: np.ndarray, dur: Any) -> np.ndarray:
    """Advance the state under a constant voltage for duration `dur` (seconds, or update ticks for discrete
    models). All updates are exact or closed-form for a constant voltage."""
    p = inst.params
    fam = inst.family
    if fam == "LD":
        ron, roff = p["R_on"], p["R_off"]
        c = p["mu_v"] * ron / (p["D"] ** 2)
        a = 0.5 * (ron - roff)
        b = roff
        f0 = b * x + a * x * x
        ft = np.clip(f0 + c * v * dur, 0.0, b + a)
        xn = (-b + np.sqrt(np.maximum(b * b + 4.0 * a * ft, 0.0))) / (2.0 * a)
        return np.clip(xn, 0.0, 1.0)
    if fam == "VT":
        pos = v >= p["v_off"]
        neg = v <= p["v_on"]
        cpos = np.where(pos, p["k_off"] * np.maximum(v / p["v_off"] - 1.0, 0.0) ** p["alpha_off"], 0.0)
        cneg = np.where(neg, abs(p["k_on"]) * np.maximum(v / p["v_on"] - 1.0, 0.0) ** p["alpha_on"], 0.0)
        z0 = np.where(neg, 1.0 - x, x)
        z0c = np.clip(z0, 0.0, 1.0 - 1e-12)
        c = cpos + cneg
        target = 0.5 * (np.arctanh(z0c) + np.arctan(z0c)) + c * dur
        z = np.where(c > 0, _vt_solve(target, z0c), z0)
        return np.clip(np.where(neg, 1.0 - z, z), 0.0, 1.0)
    if fam == "VV":
        age = aux["age"]
        write = v >= p["v_th"]
        decay = v <= p["v_hold"]
        dxdt = p["k"] * np.maximum(v / p["v_th"] - 1.0, 0.0) ** p["alpha"]
        tau, beta = max(p["tau"], EPS), p["beta"]
        t1 = age + dur
        inc = (t1 / tau) ** beta - (age / tau) ** beta
        xn = np.where(write, x + dur * dxdt, np.where(decay, x * np.exp(-inc), x))
        aux["age"] = np.where(decay, t1, 0.0)
        return np.clip(xn, 0.0, 1.0)
    if fam == "ZH":
        av = p["alpha"] * np.abs(v)
        U = av / (av + 1.0)
        B = (v > 0.0).astype(np.float64)
        a = p["r"] - U
        c = (1.0 - p["r"]) * p["G0"] + U * B
        n = np.asarray(dur, dtype=np.float64)
        an = np.power(a, n)
        xn = an * x + c * (1.0 - an) / (1.0 - a)
        return np.clip(xn, 0.0, 1.0)
    if fam == "DF":
        lam = aux["lam"]
        gp = 1.0 / (1.0 + np.exp(-p["alpha_plus"] * (v - p["delta_plus"])))
        gm = 1.0 / (1.0 + np.exp(-p["alpha_minus"] * (v + p["delta_minus"])))
        lam = np.clip(np.minimum(gm, np.maximum(lam, gp)), 0.0, 1.0)
        aux["lam"] = lam
        tau = np.maximum(p["tau0"] * np.exp(-np.abs(v) / max(p["v0"], EPS)), EPS)
        xn = lam + (x - lam) * np.exp(-dur / tau)
        return np.clip(xn, 0.0, 1.0)
    if fam == "DN":
        xmin, xmax = p["x_min"], p["x_max"]
        n = np.asarray(dur, dtype=np.float64)
        write = v > p["V_th"]
        inc = p["t_pulse"] * p["lambda"] * np.sinh(p["eta"] * np.maximum(v, 0.0))
        # Exact solution of dx/dn = inc (1 - exp(3 (x - x_max))), the continuum limit of the published pulse
        # map (per-pulse increment below 4e-4, so the step error is negligible; verified in the self-test).
        s0 = np.clip(np.exp(3.0 * (x - xmax)), 1e-300, 1.0 - 1e-15)
        K = np.log(s0) - np.log1p(-s0) + 3.0 * inc * n
        s = 1.0 / (1.0 + np.exp(-np.clip(K, -700.0, 700.0)))
        xw = xmax + np.log(np.maximum(s, 1e-300)) / 3.0
        xl = xmin + (x - xmin) * np.exp(-n * p["t_pulse"] / max(p["tau"], EPS))
        return np.clip(np.where(write, xw, xl), xmin, xmax)
    if fam in TX_LIKE:
        # Three parallel volatile state components, each with exact constant-voltage solutions:
        # saturating write above threshold, exponential relaxation otherwise.
        write = v > p["v_th"]
        r = np.maximum(v - p["v_th"], 0.0) ** p["alpha"]
        tot = np.zeros(np.broadcast(x, v).shape)
        for j in range(3):
            y = aux[f"y{j}"]
            m, kj, tj = p[f"m{j}"], p[f"k{j}"], max(p[f"tau{j}"], EPS)
            yn = np.where(write, m + (y - m) * np.exp(-kj * r * dur / m), y * np.exp(-np.asarray(dur, dtype=np.float64) / tj))
            aux[f"y{j}"] = yn
            tot = tot + yn
        return tot
    raise ValueError(fam)


def observable(inst: Instance, x: np.ndarray) -> np.ndarray:
    """Measurable output: log10 of the read conductance at a small read voltage."""
    p = inst.params
    fam = inst.family
    if fam == "LD":
        r = p["R_on"] * x + p["R_off"] * (1.0 - x)
        return -np.log10(r)
    if fam == "VT":
        r = p["R_on"] + (p["R_off"] - p["R_on"]) * x
        return -np.log10(r)
    if fam == "VV":
        r = p["R_off"] + (p["R_on"] - p["R_off"]) * x
        return -np.log10(np.maximum(r, 1.0))
    if fam == "ZH":
        return np.log10(np.maximum(x, 1e-4))
    if fam == "DF":
        r = p["R_on"] * x + p["R_off"] * (1.0 - x)
        return -np.log10(r)
    if fam == "DN":
        vr = READ_VOLTAGE
        i = (1.0 - x) * p["alpha_i"] * (1.0 - np.exp(-p["beta_i"] * vr)) + x * p["gamma_i"] * np.sinh(p["delta_i"] * vr)
        return np.log10(np.maximum(np.abs(i) / vr, 1e-30))
    if fam in TX_LIKE:
        return np.log10(1.0 + np.maximum(x, 0.0)) + p.get("log_g0", 0.0)
    raise ValueError(fam)


def norm_state(fam: str, x: np.ndarray) -> np.ndarray:
    f = FAMILIES[fam]
    return (x - f.x_lo) / (f.x_hi - f.x_lo)


# =============================================================================
# Task-blind interface calibration (datasheet quantities from nominal parameters)
# =============================================================================


def calibrate_switch_time(fam: str, max_steps: int = 300_000) -> float:
    inst = nominal_instance(fam)
    x, aux = init_state(inst, (1, 1))
    v = u_to_voltage(fam, np.ones((1, 1)))
    x0 = float(norm_state(fam, x)[0, 0])
    dt = FAMILIES[fam].dt_hint
    for step in range(1, max_steps + 1):
        x = advance(inst, x, aux, v, dt)
        if abs(float(norm_state(fam, x)[0, 0]) - x0) >= 0.35:
            return float(step if FAMILIES[fam].discrete else step * dt)
    raise RuntimeError(f"{fam} did not reach the 0.35 calibration excursion.")


def recovery_duration(fam: str) -> float:
    p = FAMILIES[fam].params
    q = RETENTION_PER_SYMBOL
    if fam == "VV":
        return p["tau"] * (-math.log(q)) ** (1.0 / p["beta"])
    if fam == "ZH":
        return float(max(1, round(math.log(q) / math.log(p["r"]))))
    if fam == "DN":
        return float(max(1, round(math.log(q) / (-p["t_pulse"] / p["tau"]))))
    if fam in TX_LIKE:
        j = int(np.argmax([p[f"m{i}"] for i in range(3)]))
        return -p[f"tau{j}"] * math.log(q)
    return 0.0


T_SWITCH: Dict[str, float] = {}
T_RECOVER: Dict[str, float] = {}
OP_GAIN_IDX = np.repeat(np.arange(len(GAINS)), len(FRACS))
OP_FRAC_IDX = np.tile(np.arange(len(FRACS)), len(GAINS))
N_OPS = len(OP_GAIN_IDX)


def op_duration(fam: str, op: int) -> float:
    d = FRACS[OP_FRAC_IDX[op]] * T_SWITCH[fam]
    if FAMILIES[fam].discrete:
        return float(max(1, int(round(d))))
    return float(d)


def op_key(fam: str, op: int) -> Tuple[int, float]:
    return int(OP_GAIN_IDX[op]), float(np.round(op_duration(fam, op), 15))


# =============================================================================
# Simulation
# =============================================================================


def bin_edges(L: int) -> np.ndarray:
    return np.linspace(0, L, N_BINS + 1).round().astype(int)


def features_from_traj(traj: np.ndarray) -> np.ndarray:
    """traj (n, L) -> (n, 2 * N_BINS): bin means and bin standard deviations."""
    traj = np.asarray(traj, dtype=np.float64)
    e = bin_edges(traj.shape[1])
    cnt = np.diff(e).astype(np.float64)
    s = np.add.reduceat(traj, e[:-1], axis=1)
    q = np.add.reduceat(traj * traj, e[:-1], axis=1)
    mu = s / cnt
    sd = np.sqrt(np.maximum(q / cnt - mu * mu, 0.0))
    return np.concatenate([mu, sd], axis=1).astype(np.float32)


def simulate(inst: Instance, U: np.ndarray, ops: Sequence[int], mode: str, key: Tuple[Any, ...]) -> np.ndarray:
    """Drive the device with each row of U (one independent sequence per row) at each operating point.

    mode "features": returns (len(ops), n, 2*N_BINS) float32
    mode "traj":     returns (n, L, len(ops)) float32 trajectories of the noisy observable

    Noise streams are keyed by the instance's effective physical configuration, so the same operating point
    produces identical noise regardless of which other operating points are simulated alongside it.
    """
    fam = inst.family
    U = np.asarray(U, dtype=np.float64)
    n, L = U.shape
    ops = list(ops)
    keys = [op_key(fam, o) for o in ops]
    uniq = sorted(set(keys))
    col_of = {k: i for i, k in enumerate(uniq)}
    m = len(uniq)
    gains = np.array([GAINS[k[0]] for k in uniq])[None, :]
    durs = np.array([k[1] for k in uniq])[None, :]
    rngs = [np.random.default_rng(stable_seed(*key, inst.iid, "cfg", k[0], repr(k[1]))) for k in uniq]
    x, aux = init_state(inst, (n, m))
    v_rest = float(u_to_voltage(fam, np.array(rest_input(fam))))
    rec = T_RECOVER[fam]
    lo, hi = FAMILIES[fam].x_lo, FAMILIES[fam].x_hi
    span = hi - lo
    if fam == "DN":
        lo, hi = inst.params["x_min"], inst.params["x_max"]
    nc = np.empty((n, m), dtype=np.float32)
    nr = np.empty((n, m), dtype=np.float32)
    if mode == "features":
        e = bin_edges(L)
        bin_of = np.searchsorted(e, np.arange(L), side="right") - 1
        S = np.zeros((n, m, N_BINS))
        Q = np.zeros((n, m, N_BINS))
    else:
        out = np.empty((n, L, m), dtype=np.float32)
    for t in range(L):
        v = u_to_voltage(fam, gains * U[:, t:t + 1])
        x = advance(inst, x, aux, v, durs)
        if SIGMA_C2C > 0 or SIGMA_READ > 0:
            for j, r in enumerate(rngs):
                nc[:, j] = r.standard_normal(n, dtype=np.float32)
                nr[:, j] = r.standard_normal(n, dtype=np.float32)
        if SIGMA_C2C > 0:
            x_old = x
            x = np.clip(x + SIGMA_C2C * span * nc, lo, hi)
            if fam in TX_LIKE:
                ratio = np.where(x_old > 1e-12, x / np.maximum(x_old, 1e-12), 1.0)
                for j in range(3):
                    aux[f"y{j}"] = aux[f"y{j}"] * ratio
        o = observable(inst, x)
        if SIGMA_READ > 0:
            o = o + SIGMA_READ * nr
        if mode == "features":
            b = bin_of[t]
            S[:, :, b] += o
            Q[:, :, b] += o * o
        else:
            out[:, t, :] = o
        if rec > 0:
            x = advance(inst, x, aux, np.full_like(x, v_rest), rec)
    idx = [col_of[k] for k in keys]
    if mode == "features":
        cnt = np.diff(bin_edges(L)).astype(np.float64)
        mu = S / cnt
        sd = np.sqrt(np.maximum(Q / cnt - mu * mu, 0.0))
        feats = np.concatenate([mu, sd], axis=2).astype(np.float32)   # (n, m, F)
        return np.transpose(feats, (1, 0, 2))[idx]
    return out[:, :, idx]


def rest_observable(inst: Instance) -> float:
    x, _ = init_state(inst, (1,))
    return float(observable(inst, x)[0])


# =============================================================================
# Self-tests of the closed-form updates
# =============================================================================


def self_tests() -> pd.DataFrame:
    rows = []
    # Daniels: closed form versus the published per-pulse map.
    inst = nominal_instance("DN")
    p = inst.params
    for v in [0.62, 0.7, 0.8]:
        x_loop = p["x0"]
        inc = p["t_pulse"] * p["lambda"] * math.sinh(p["eta"] * v)
        for _ in range(2000):
            x_loop = min(max(x_loop + inc * (1 - math.exp(3 * x_loop) / math.exp(3 * p["x_max"])), p["x_min"]), p["x_max"])
        x_cf = float(advance(inst, np.array([p["x0"]]), {}, np.array([v]), 2000.0)[0])
        rows.append({"test": f"DN closed form vs pulse map, V={v}", "abs_error": abs(x_cf - x_loop)})
    # VTEAM: exact inverse versus fine RK4.
    inst = nominal_instance("VT")
    for v, dur in [(0.5, 0.05), (-0.5, 0.002), (-0.4, 0.01)]:
        x_cf = float(advance(inst, np.array([0.3]), {}, np.array([v]), dur)[0])
        x = 0.3
        pp = inst.params
        def f(z):
            if v >= pp["v_off"]:
                return pp["k_off"] * (v / pp["v_off"] - 1) ** pp["alpha_off"] * max(1 - z ** 4, 0)
            if v <= pp["v_on"]:
                return pp["k_on"] * (v / pp["v_on"] - 1) ** pp["alpha_on"] * max(1 - (z - 1) ** 4, 0)
            return 0.0
        nsub = 20000
        h = dur / nsub
        for _ in range(nsub):
            k1 = f(x); k2 = f(min(max(x + 0.5 * h * k1, 0), 1)); k3 = f(min(max(x + 0.5 * h * k2, 0), 1))
            k4 = f(min(max(x + h * k3, 0), 1))
            x = min(max(x + h / 6 * (k1 + 2 * k2 + 2 * k3 + k4), 0), 1)
        rows.append({"test": f"VT exact vs RK4, V={v}, dt={dur}", "abs_error": abs(x_cf - x)})
    # Zhong: closed-form affine iterate versus explicit iteration.
    inst = nominal_instance("ZH")
    pp = inst.params
    x = 0.5
    for _ in range(7):
        U_ = pp["alpha"] * 2.0 / (pp["alpha"] * 2.0 + 1)
        x = pp["r"] * x + (1 - pp["r"]) * pp["G0"] + U_ * (1 - x)
    x_cf = float(advance(inst, np.array([0.5]), {}, np.array([2.0]), 7.0)[0])
    rows.append({"test": "ZH closed form vs iteration, 7 ticks", "abs_error": abs(x_cf - x)})
    return pd.DataFrame(rows)


# =============================================================================
# Data
# =============================================================================


@dataclass
class Dataset:
    name: str
    tier: str
    X_train_raw: np.ndarray
    X_test_raw: np.ndarray
    U_train: np.ndarray
    U_test: np.ndarray
    y_train: np.ndarray
    y_test: np.ndarray
    n_classes: int


def ensure_aeon() -> None:
    try:
        importlib.import_module("aeon")
    except ImportError:
        log("Installing aeon ...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "aeon>=1.0"])


def robust_scale(Xtr: np.ndarray, X: np.ndarray) -> np.ndarray:
    c = float(np.median(Xtr))
    s = float(np.quantile(np.abs(Xtr - c), 0.995))
    if not np.isfinite(s) or s <= EPS:
        s = float(np.std(Xtr)) or 1.0
    return np.clip((X - c) / s, -1.0, 1.0)


def synthetic_dataset(name: str, seed: int) -> Dataset:
    rng = np.random.default_rng(seed)
    L = 96
    t = np.linspace(0, 1, L)
    def make(n):
        y = rng.integers(0, 3, n)
        X = np.empty((n, L))
        for i, c in enumerate(y):
            ph = rng.uniform(0, 2 * np.pi)
            base = [np.sin(2 * np.pi * 3 * t + ph), np.sign(np.sin(2 * np.pi * 3 * t + ph)),
                    np.sin(2 * np.pi * 6 * t + ph) * t][c]
            X[i] = base + 0.3 * rng.standard_normal(L)
        return X, y
    Xtr, ytr = make(60)
    Xte, yte = make(60)
    return Dataset(name, "smoke", Xtr, Xte, robust_scale(Xtr, Xtr), robust_scale(Xtr, Xte), ytr, yte, 3)


def load_dataset(name: str) -> Dataset:
    ensure_aeon()
    from aeon.datasets import load_classification
    Xtr, ytr = load_classification(name=name, split="train", extract_path=str(DATA_DIR))
    Xte, yte = load_classification(name=name, split="test", extract_path=str(DATA_DIR))
    Xtr = np.asarray(Xtr, dtype=np.float64)
    Xte = np.asarray(Xte, dtype=np.float64)
    if Xtr.ndim != 3 or Xtr.shape[1] != 1:
        raise RuntimeError(f"{name}: expected univariate equal-length data, got {Xtr.shape}")
    Xtr, Xte = Xtr[:, 0, :], Xte[:, 0, :]
    if not (np.all(np.isfinite(Xtr)) and np.all(np.isfinite(Xte))):
        raise RuntimeError(f"{name}: non-finite values")
    classes = np.unique(np.asarray(ytr).astype(str))
    cmap = {c: i for i, c in enumerate(classes)}
    ytr_i = np.array([cmap[v] for v in np.asarray(ytr).astype(str)])
    yte_i = np.array([cmap[v] for v in np.asarray(yte).astype(str)])
    tier = "prior" if name in PRIOR_DATASETS else "fresh"
    return Dataset(name, tier, Xtr, Xte, robust_scale(Xtr, Xtr), robust_scale(Xtr, Xte), ytr_i, yte_i, len(classes))


# =============================================================================
# Readout, metrics
# =============================================================================


def balanced_accuracy(y: np.ndarray, yp: np.ndarray, C: int) -> float:
    rec = []
    for c in range(C):
        mask = y == c
        if mask.any():
            rec.append(float(np.mean(yp[mask] == c)))
    return float(np.mean(rec))


def nba(ba: float, C: int) -> float:
    return (ba - 1.0 / C) / (1.0 - 1.0 / C)


def std_stats(F: np.ndarray) -> Tuple[np.ndarray, np.ndarray]:
    mu = F.mean(axis=0).astype(np.float64)
    sd = F.std(axis=0).astype(np.float64)
    sd[sd < 1e-9] = 1.0
    return mu, sd


def class_weights(y: np.ndarray, C: int) -> np.ndarray:
    counts = np.bincount(y, minlength=C).astype(np.float64)
    present = int((counts > 0).sum())
    return len(y) / (present * counts[y])


def _gram(Z: np.ndarray, y: np.ndarray, C: int):
    w = class_weights(y, C)
    Y = -np.ones((len(y), C))
    Y[np.arange(len(y)), y] = 1.0
    ws = w.sum()
    mz = (w @ Z) / ws
    my = (w @ Y) / ws
    Zc = Z - mz
    Zw = Zc * w[:, None]
    return Zw.T @ Zc, Zw.T @ (Y - my), mz, my, ws


@dataclass
class Readout:
    W: np.ndarray
    b: np.ndarray

    def predict(self, Z: np.ndarray) -> np.ndarray:
        return np.argmax(Z @ self.W + self.b, axis=1)


def ridge_fit(Z: np.ndarray, y: np.ndarray, C: int, alpha_abs: float,
              prior: Optional[np.ndarray] = None) -> Readout:
    A, B, mz, my, _ = _gram(np.asarray(Z, dtype=np.float64), y, C)
    reg = alpha_abs * np.eye(A.shape[0])
    rhs = B if prior is None else B + alpha_abs * prior
    try:
        W = np.linalg.solve(A + reg, rhs)
    except np.linalg.LinAlgError:
        W = np.linalg.lstsq(A + reg, rhs, rcond=None)[0]
    return Readout(W, my - mz @ W)


def ridge_multi_alpha(Z: np.ndarray, y: np.ndarray, C: int, alphas_rel: np.ndarray) -> List[Readout]:
    A, B, mz, my, ws = _gram(np.asarray(Z, dtype=np.float64), y, C)
    e, V = np.linalg.eigh(A)
    VB = V.T @ B
    out = []
    for a in alphas_rel:
        W = V @ (VB / (e + a * ws)[:, None])
        out.append(Readout(W, my - mz @ W))
    return out


def make_folds(y: np.ndarray, seed: int) -> List[Tuple[np.ndarray, np.ndarray]]:
    min_count = int(np.bincount(y).min())
    k = int(min(3, max(2, min_count)))
    skf = StratifiedKFold(n_splits=k, shuffle=True, random_state=seed % (2**31))
    return list(skf.split(np.zeros(len(y)), y))


def select_and_test(Ftr: np.ndarray, Fte: np.ndarray, ytr: np.ndarray, yte: np.ndarray, C: int,
                    folds) -> Tuple[float, int, float]:
    """CV over the relative ridge grid; returns (cv_ba, alpha_index, test_ba)."""
    scores = np.zeros(len(ALPHA_REL))
    for tr, va in folds:
        mu, sd = std_stats(Ftr[tr])
        Ztr = (Ftr[tr] - mu) / sd
        Zva = (Ftr[va] - mu) / sd
        for i, ro in enumerate(ridge_multi_alpha(Ztr, ytr[tr], C, ALPHA_REL)):
            scores[i] += balanced_accuracy(ytr[va], ro.predict(Zva), C)
    scores /= len(folds)
    ai = int(np.argmax(scores))
    mu, sd = std_stats(Ftr)
    ro = ridge_multi_alpha((Ftr - mu) / sd, ytr, C, ALPHA_REL[ai:ai + 1])[0]
    return float(scores[ai]), ai, balanced_accuracy(yte, ro.predict((Fte - mu) / sd), C)


# =============================================================================
# Alignment maps
# =============================================================================


def lag_design(r: np.ndarray, taps: int, pad: float, quadratic: bool) -> np.ndarray:
    """r (n, L) -> (n*L, P) causal lag design, padded with the device rest value."""
    n, L = r.shape
    rp = np.concatenate([np.full((n, taps - 1), pad, dtype=np.float64), r.astype(np.float64)], axis=1)
    cols = [rp[:, taps - 1 - k: taps - 1 - k + L] for k in range(taps)]
    X = np.stack(cols, axis=2).reshape(n * L, taps)
    if quadratic:
        X = np.concatenate([X, X * X], axis=1)
    return X


@dataclass
class LagMap:
    taps: int
    quadratic: bool
    pad: float
    mx: np.ndarray
    sx: np.ndarray
    beta: np.ndarray
    my: float

    def apply(self, r: np.ndarray) -> np.ndarray:
        n, L = r.shape
        out = np.full((n, L), self.my, dtype=np.float64)
        rp = np.concatenate([np.full((n, self.taps - 1), self.pad), r.astype(np.float64)], axis=1)
        P = self.taps
        for k in range(self.taps):
            lagk = rp[:, P - 1 - k: P - 1 - k + L]
            out += self.beta[k] * (lagk - self.mx[k]) / self.sx[k]
            if self.quadratic:
                out += self.beta[P + k] * (lagk * lagk - self.mx[P + k]) / self.sx[P + k]
        return out.astype(np.float32)


def fit_lag_map(r_src: np.ndarray, r_tgt: np.ndarray, pad: float, taps: int, quadratic: bool) -> LagMap:
    X = lag_design(r_tgt, taps, pad, quadratic)
    y = r_src.reshape(-1).astype(np.float64)
    mx = X.mean(0)
    sx = X.std(0)
    sx[sx < 1e-12] = 1.0
    Xs = (X - mx) / sx
    my = float(y.mean())
    A = Xs.T @ Xs
    lam = 1e-3 * np.trace(A) / A.shape[0] + 1e-12
    beta = np.linalg.solve(A + lam * np.eye(A.shape[0]), Xs.T @ (y - my))
    return LagMap(taps, quadratic, pad, mx, sx, beta, my)


def r2_score_flat(y: np.ndarray, yhat: np.ndarray) -> float:
    y = np.asarray(y, dtype=np.float64).reshape(-1)
    yhat = np.asarray(yhat, dtype=np.float64).reshape(-1)
    ss = np.sum((y - y.mean()) ** 2)
    if ss < 1e-20:
        return float("nan")
    return float(1.0 - np.sum((y - yhat) ** 2) / ss)


def coral_matrix(Zs: np.ndarray, Zt: np.ndarray) -> np.ndarray:
    def msqrt(C, inv):
        e, V = np.linalg.eigh(C)
        e = np.maximum(e, 1e-10)
        return (V * (e ** (-0.5 if inv else 0.5))) @ V.T
    Cs = np.cov(Zs, rowvar=False) + CORAL_EPS * np.eye(Zs.shape[1])
    Ct = np.cov(Zt, rowvar=False) + CORAL_EPS * np.eye(Zt.shape[1])
    return msqrt(Ct, True) @ msqrt(Cs, False)


# =============================================================================
# Capacity
# =============================================================================


def legendre(k: int, u: np.ndarray) -> np.ndarray:
    if k == 1:
        return u
    if k == 2:
        return 0.5 * (3 * u * u - 1)
    if k == 3:
        return 0.5 * (5 * u ** 3 - 3 * u)
    raise ValueError(k)


def capacity_targets(Uprobe: np.ndarray, t0: int) -> Tuple[np.ndarray, np.ndarray]:
    """Return (targets (rows, n_targets), degree per target) on rows t >= t0 of every episode."""
    E, L = Uprobe.shape
    def lagged(d):
        return Uprobe[:, t0 - d: L - d].reshape(-1)
    cols, deg = [], []
    for d in range(CAP_MAX_DELAY_D1 + 1):
        cols.append(lagged(d)); deg.append(1)
    D2 = CAP_MAX_DELAY_D2
    for i in range(D2 + 1):
        cols.append(legendre(2, lagged(i))); deg.append(2)
        for j in range(i + 1, D2 + 1):
            cols.append(lagged(i) * lagged(j)); deg.append(2)
    D3 = CAP_MAX_DELAY_D3
    for i in range(D3 + 1):
        cols.append(legendre(3, lagged(i))); deg.append(3)
        for j in range(D3 + 1):
            if j != i:
                cols.append(legendre(2, lagged(i)) * lagged(j)); deg.append(3)
        for j in range(i + 1, D3 + 1):
            for k in range(j + 1, D3 + 1):
                cols.append(lagged(i) * lagged(j) * lagged(k)); deg.append(3)
    Y = np.stack(cols, axis=1)
    return Y - Y.mean(0), np.array(deg)


def capacities(state_signal: np.ndarray, Yc: np.ndarray, deg: np.ndarray, t0: int) -> Dict[str, float]:
    E, L = state_signal.shape
    X = np.stack([state_signal[:, t0 - k: L - k].reshape(-1) for k in range(CAP_STATE_LAGS)], axis=1)
    X = X - X.mean(0)
    sc = X.std(0)
    keep = sc > 1e-12
    if not keep.any():
        return {"C1": 0.0, "C2": 0.0, "C3": 0.0, "C_total": 0.0, "linear_fraction": float("nan")}
    Q, R = np.linalg.qr(X[:, keep] / sc[keep])
    diag = np.abs(np.diag(R))
    Q = Q[:, diag > 1e-8 * diag.max()]
    proj = Q.T @ Yc
    cap = np.sum(proj * proj, axis=0) / np.sum(Yc * Yc, axis=0)
    cap[cap < max(CAP_THRESHOLD, 4.0 * (CAP_STATE_LAGS + 1) / X.shape[0])] = 0.0
    c1, c2, c3 = (float(cap[deg == d].sum()) for d in (1, 2, 3))
    tot = c1 + c2 + c3
    return {"C1": c1, "C2": c2, "C3": c3, "C_total": tot,
            "linear_fraction": c1 / tot if tot > 1e-9 else float("nan")}


# =============================================================================
# Statistics helpers
# =============================================================================


def wilcoxon_less(x: np.ndarray, y: np.ndarray) -> Tuple[float, int]:
    """One-sided p for median(x - y) < 0."""
    d = np.asarray(x, float) - np.asarray(y, float)
    d = d[np.isfinite(d)]
    if len(d) < 3 or np.all(np.abs(d) < 1e-12):
        return float("nan"), len(d)
    return float(wilcoxon(d, alternative="less", zero_method="wilcox").pvalue), len(d)


def holm(pvals: Dict[str, float]) -> Dict[str, float]:
    items = [(k, v) for k, v in pvals.items() if np.isfinite(v)]
    items.sort(key=lambda kv: kv[1])
    m = len(items)
    out, running = {}, 0.0
    for i, (k, p) in enumerate(items):
        running = max(running, min(1.0, (m - i) * p))
        out[k] = running
    for k, v in pvals.items():
        out.setdefault(k, float("nan"))
    return out


def boot_median_ci(x: np.ndarray, seed: int) -> Tuple[float, float, float]:
    x = np.asarray(x, float)
    x = x[np.isfinite(x)]
    if len(x) == 0:
        return (float("nan"),) * 3
    rng = np.random.default_rng(seed)
    meds = np.median(rng.choice(x, size=(N_BOOT, len(x)), replace=True), axis=1)
    return float(np.median(x)), float(np.quantile(meds, 0.025)), float(np.quantile(meds, 0.975))


def hier_boot_ci(df: pd.DataFrame, col: str, seed: int) -> Tuple[float, float, float]:
    """Resample datasets, then pairs within datasets; statistic is the median of dataset medians."""
    groups = [g[col].dropna().to_numpy() for _, g in df.groupby("dataset")]
    groups = [g for g in groups if len(g)]
    if not groups:
        return (float("nan"),) * 3
    obs = float(np.median([np.median(g) for g in groups]))
    rng = np.random.default_rng(seed)
    stats = np.empty(N_BOOT)
    for b in range(N_BOOT):
        pick = rng.integers(0, len(groups), len(groups))
        stats[b] = np.median([np.median(rng.choice(groups[i], len(groups[i]))) for i in pick])
    return obs, float(np.quantile(stats, 0.025)), float(np.quantile(stats, 0.975))


# =============================================================================
# Figure helpers (all spines, labelled bounds at tick multiples, outlined legends)
# =============================================================================

SPINE_W = 1.6
LEGEND_W = 0.8
plt.rcParams.update({
    "font.family": "sans-serif", "font.sans-serif": ["Arial", "DejaVu Sans"],
    "font.size": 15, "axes.labelsize": 17, "xtick.labelsize": 14, "ytick.labelsize": 14,
    "legend.fontsize": 13, "axes.linewidth": SPINE_W, "xtick.major.width": SPINE_W,
    "ytick.major.width": SPINE_W, "savefig.dpi": 600, "savefig.bbox": "tight", "figure.dpi": 110,
})
FAM_COLORS = {"LD": "#1b6ca8", "VT": "#e08a1e", "VV": "#2a9d5c", "ZH": "#c0392b", "DF": "#7b4fa0", "DN": "#8c6d3f"}


def _nice_step(span: float, n: int = 5) -> float:
    span = max(span, 1e-12)
    raw = span / n
    mag = 10 ** math.floor(math.log10(raw))
    for mult in (1, 2, 2.5, 5, 10):
        if mult * mag >= raw:
            return mult * mag
    return 10 * mag


def _decimals(step: float) -> int:
    d = 0
    while d < 6 and abs(round(step * 10 ** d) - step * 10 ** d) > 1e-6:
        d += 1
    return d


def set_bounds(ax, axis: str, lo: float, hi: float, step: Optional[float] = None, n: int = 5) -> None:
    if not (np.isfinite(lo) and np.isfinite(hi)):
        lo, hi = 0.0, 1.0
    if hi - lo < 1e-12:
        hi = lo + 1.0
    step = step or _nice_step(hi - lo, n)
    a = math.floor(lo / step + 1e-9) * step
    b = math.ceil(hi / step - 1e-9) * step
    if b <= a:
        b = a + step
    ticks = np.arange(a, b + step / 2, step)
    d = _decimals(step)
    fmt = FuncFormatter(lambda v, _: (f"{v:.{d}f}" if abs(v) > 10 ** (-d - 1) else f"{0:.{d}f}"))
    if axis == "y":
        ax.set_ylim(a, b); ax.yaxis.set_major_locator(FixedLocator(ticks)); ax.yaxis.set_major_formatter(fmt)
    else:
        ax.set_xlim(a, b); ax.xaxis.set_major_locator(FixedLocator(ticks)); ax.xaxis.set_major_formatter(fmt)


def style_ax(ax) -> None:
    for s in ax.spines.values():
        s.set_visible(True); s.set_linewidth(SPINE_W)


def outlined_legend(ax, **kw):
    leg = ax.legend(frameon=True, fancybox=False, edgecolor="black", framealpha=1.0, **kw)
    leg.get_frame().set_linewidth(LEGEND_W)
    return leg


def panel_letters(fig, axes: Sequence[Any]) -> None:
    fig.canvas.draw()
    rend = fig.canvas.get_renderer()
    inv = fig.transFigure.inverted()
    for i, ax in enumerate(axes):
        bb = ax.yaxis.label.get_window_extent(rend)
        x0 = inv.transform((bb.x0, 0))[0]
        top = inv.transform((0, ax.get_window_extent(rend).y1))[1]
        fig.text(x0, top + 0.015, chr(ord("a") + i), fontsize=20, fontweight="bold", ha="left", va="bottom")


def save_fig(fig, stem: str) -> None:
    fig.savefig(DIRS["figures"] / f"{stem}.png")
    fig.savefig(DIRS["figures"] / f"{stem}.pdf")
    plt.close(fig)





# =============================================================================
# v4: preregistered confirmatory study and reviewer-question battery
# =============================================================================

from contextlib import contextmanager
from sklearn.linear_model import LogisticRegression
from sklearn.model_selection import KFold

V4_SEEDS = [44, 55, 66] if not SMOKE else [44]
ROBUST_SEED = V4_SEEDS[0]
HELD_OUT_DATASETS = ["MoteStrain", "SonyAIBORobotSurface1", "Coffee", "Meat", "Lightning7",
                     "MedicalImages", "ProximalPhalanxOutlineCorrect", "Ham", "OSULeaf", "Fish"]
V4_D2D = [0.05, 0.20, 0.50] if not SMOKE else [0.05, 0.50]
BUDGETS = [8, 16, 32, 64, 128] if not SMOKE else [8, 16]
K_V4 = [1, 5] if not SMOKE else [1, 2]
K_MAIN = K_V4[-1]
RFF_DIM = 512 if not SMOKE else 128
N_ALT = 256 if not SMOKE else 48
MAP_ALPHAS_LIN = [1e-4, 1e-3, 1e-2, 1e-1, 1.0]
MAP_ALPHAS_RFF = [1e-5, 1e-4, 1e-3, 1e-2, 1e-1]
LR_CS = [0.1, 1.0, 10.0]
NOMINAL_NOISE = (SIGMA_C2C, SIGMA_READ)
NOISE_LEVELS = {"noise_none": (0.0, 0.0), "noise_high": (3.0 * 0.002, 3.0 * 0.01)}
FEATURE_VARIANTS = [("bins12", 12, "ms"), ("bins48", 48, "ms"), ("means24", 24, "m")]
ROBUST_ORDER = [("noise_none", "No noise"), ("main", "Nominal"), ("noise_high", "High noise"), ("bins12", "12 bins"),
                ("bins48", "48 bins"), ("means24", "Means only"), ("logistic", "Logistic")]
CAP_T0 = max(CAP_MAX_DELAY_D1, CAP_STATE_LAGS)
V3_NAME = "portability_v3_diagnostics" + ("_ideal" if IDEAL else "") + ("_smoke" if SMOKE else "")

LADDER = [("zero_shot", "Zero\nshot"), ("moment", "Moment\nmatching"), ("sign_moment", "Sign\ncorrected"),
          ("diag_paired", "Diagonal\npaired"), ("paired_linear", "Paired\nlinear"), ("paired_rff", "Paired\nnonlinear"),
          ("distill_rff", "Score\ndistillation"), ("pseudo_label", "Pseudo\nlabels"), ("twin_paired_rff", "Twin\nnonlinear")]


@contextmanager
def noise_level(c2c: float, rd: float):
    global SIGMA_C2C, SIGMA_READ
    old = (SIGMA_C2C, SIGMA_READ)
    SIGMA_C2C, SIGMA_READ = c2c, rd
    try:
        yield
    finally:
        SIGMA_C2C, SIGMA_READ = old


def probe_inputs(seed: int) -> np.ndarray:
    rng = np.random.default_rng(stable_seed("probe_input", seed, MASTER_SEED))
    return rng.uniform(-1.0, 1.0, size=(PROBE_EPISODES, PROBE_LENGTH))


def white_inputs(n: int, L: int, seed: int) -> np.ndarray:
    return np.random.default_rng(seed).uniform(-1.0, 1.0, size=(n, L))


def matched_inputs(U: np.ndarray, n: int, L: int, seed: int) -> np.ndarray:
    """Gaussian AR(1) sequences matched to the lag-1 autocorrelation of U, quantile-mapped to the marginal of U."""
    a, b = U[:, :-1].reshape(-1), U[:, 1:].reshape(-1)
    rho = float(np.clip(np.corrcoef(a, b)[0, 1], -0.999, 0.999)) if np.std(a) > 0 else 0.0
    rng = np.random.default_rng(seed)
    g = np.empty((n, L))
    g[:, 0] = rng.standard_normal(n)
    s = math.sqrt(max(1.0 - rho * rho, 1e-6))
    for t in range(1, L):
        g[:, t] = rho * g[:, t - 1] + s * rng.standard_normal(n)
    q = np.sort(U.reshape(-1))
    ranks = np.argsort(np.argsort(g.reshape(-1)))
    pos = (ranks + 0.5) / len(ranks) * (len(q) - 1)
    return q[pos.round().astype(int)].reshape(n, L)


def generic_inputs(pool: Dict[str, np.ndarray], exclude: str, n: int, L: int, seed: int) -> np.ndarray:
    """Unlabeled sequences from other datasets, linearly resampled to length L."""
    names = [k for k in pool if k != exclude]
    rng = np.random.default_rng(seed)
    out = np.empty((n, L))
    grid = np.linspace(0, 1, L)
    for i in range(n):
        U = pool[names[rng.integers(len(names))]]
        row = U[rng.integers(len(U))]
        out[i] = np.interp(grid, np.linspace(0, 1, len(row)), row)
    return out


def features_variant(traj: np.ndarray, nb: int, stats: str) -> np.ndarray:
    traj = np.asarray(traj, dtype=np.float64)
    e = np.linspace(0, traj.shape[1], nb + 1).round().astype(int)
    cnt = np.diff(e).astype(np.float64)
    mu = np.add.reduceat(traj, e[:-1], axis=1) / cnt
    if stats == "m":
        return mu.astype(np.float32)
    q = np.add.reduceat(traj * traj, e[:-1], axis=1)
    sd = np.sqrt(np.maximum(q / cnt - mu * mu, 0.0))
    return np.concatenate([mu, sd], axis=1).astype(np.float32)


def column_corr(A: np.ndarray, B: np.ndarray) -> np.ndarray:
    A = np.asarray(A, dtype=np.float64)
    B = np.asarray(B, dtype=np.float64)
    sa, sb = A.std(0), B.std(0)
    r = ((A - A.mean(0)) * (B - B.mean(0))).mean(0) / np.maximum(sa * sb, 1e-24)
    r[(sa < 1e-9) | (sb < 1e-9)] = np.nan
    return r


def linear_cka(X: np.ndarray, Y: np.ndarray) -> float:
    X = X - X.mean(0)
    Y = Y - Y.mean(0)
    num = np.linalg.norm(X.T @ Y, "fro") ** 2
    den = np.linalg.norm(X.T @ X, "fro") * np.linalg.norm(Y.T @ Y, "fro")
    return float(num / den) if den > 1e-20 else float("nan")


def feature_r2(Y: np.ndarray, Yhat: np.ndarray) -> float:
    ss = np.sum((Y - Y.mean(0)) ** 2)
    return float(1.0 - np.sum((Y - Yhat) ** 2) / ss) if ss > 1e-20 else float("nan")


def make_instance_sigma(fam: str, seed: int, role: str, sigma: float) -> Instance:
    global SIGMA_D2D
    old = SIGMA_D2D
    SIGMA_D2D = sigma
    try:
        return make_instance(fam, seed, role)
    finally:
        SIGMA_D2D = old


def parameter_distance(a: Instance, b: Instance) -> float:
    if a.family != b.family:
        return float("nan")
    vals = []
    for k, va in a.params.items():
        vb = b.params.get(k)
        if vb is None or va == 0 or vb == 0 or np.sign(va) != np.sign(vb) or k in ("x0", "lambda0"):
            continue
        vals.append(math.log(abs(vb) / abs(va)))
    return float(np.sqrt(np.mean(np.square(vals)))) if vals else 0.0


def _ridge_path(X: np.ndarray, Y: np.ndarray, lams: Sequence[float]) -> List[np.ndarray]:
    """Ridge solutions for several penalties from one eigendecomposition (primal or dual, whichever is smaller)."""
    n, p = X.shape
    if n >= p:
        e, V = np.linalg.eigh(X.T @ X)
        VB = V.T @ (X.T @ Y)
        return [V @ (VB / (e + lam)[:, None]) for lam in lams]
    e, V = np.linalg.eigh(X @ X.T)
    VY = V.T @ Y
    return [X.T @ (V @ (VY / (e + lam)[:, None])) for lam in lams]


def fit_feature_map(Zt: np.ndarray, Zs: np.ndarray, kind: str, seed: int):
    """Ridge map from target features (or their random Fourier features) to source features or scores.
    Penalty chosen by 3-fold CV mean squared error. Returns a callable."""
    rng = np.random.default_rng(seed)
    Zt = np.asarray(Zt, dtype=np.float64)
    if kind == "rff":
        d = Zt.shape[1]
        Wr = rng.standard_normal((d, RFF_DIM)) / math.sqrt(d)
        br = rng.uniform(0, 2 * math.pi, RFF_DIM)
        phi = lambda Z: math.sqrt(2.0 / RFF_DIM) * np.cos(np.asarray(Z, dtype=np.float64) @ Wr + br)
        alphas = MAP_ALPHAS_RFF
    else:
        phi = lambda Z: np.asarray(Z, dtype=np.float64)
        alphas = MAP_ALPHAS_LIN
    X = phi(Zt)
    Y = np.asarray(Zs, dtype=np.float64)
    mx, my = X.mean(0), Y.mean(0)
    Xc, Yc = X - mx, Y - my
    n = len(X)
    if n >= 6:
        fold = rng.permutation(n) % 3
        err = np.zeros(len(alphas))
        for f in range(3):
            tr, va = fold != f, fold == f
            for i, W in enumerate(_ridge_path(Xc[tr], Yc[tr], [a * tr.sum() for a in alphas])):
                err[i] += np.sum((Yc[va] - Xc[va] @ W) ** 2)
        a = alphas[int(np.argmin(err))]
    else:
        a = alphas[len(alphas) // 2]
    W = _ridge_path(Xc, Yc, [a * n])[0]
    return lambda Z: (phi(Z) - mx) @ W + my


def pseudo_label_ba(Ztr: np.ndarray, Zte: np.ndarray, yp: np.ndarray, yte: np.ndarray, C: int, seed: int) -> float:
    """Target readout trained on the source device's predicted labels; penalty chosen by agreement with them."""
    present = np.unique(yp)
    if len(present) < 2:
        return balanced_accuracy(yte, np.full(len(yte), present[0]), C)
    scores = np.zeros(len(ALPHA_REL))
    for tr, va in KFold(3, shuffle=True, random_state=seed % (2**31)).split(Ztr):
        if len(np.unique(yp[tr])) < 2:
            continue
        for i, ro in enumerate(ridge_multi_alpha(Ztr[tr], yp[tr], C, ALPHA_REL)):
            scores[i] += balanced_accuracy(yp[va], ro.predict(Ztr[va]), C)
    ai = int(np.argmax(scores))
    ro = ridge_multi_alpha(Ztr, yp, C, ALPHA_REL[ai:ai + 1])[0]
    return balanced_accuracy(yte, ro.predict(Zte), C)


def fit_logistic(Z: np.ndarray, y: np.ndarray, C: int, folds) -> LogisticRegression:
    best, bc = -1.0, LR_CS[0]
    for c in LR_CS:
        sc = 0.0
        for tr, va in folds:
            m = LogisticRegression(C=c, max_iter=3000, class_weight="balanced").fit(Z[tr], y[tr])
            sc += balanced_accuracy(y[va], m.predict(Z[va]), C)
        if sc > best:
            best, bc = sc, c
    return LogisticRegression(C=bc, max_iter=3000, class_weight="balanced").fit(Z, y)


def capacity_profile(state_signal: np.ndarray, Yc: np.ndarray, t0: int) -> np.ndarray:
    E, L = state_signal.shape
    X = np.stack([state_signal[:, t0 - k: L - k].reshape(-1) for k in range(CAP_STATE_LAGS)], axis=1)
    X = X - X.mean(0)
    sc = X.std(0)
    keep = sc > 1e-12
    if not keep.any():
        return np.zeros(Yc.shape[1])
    Q, R = np.linalg.qr(X[:, keep] / sc[keep])
    diag = np.abs(np.diag(R))
    Q = Q[:, diag > 1e-8 * diag.max()]
    proj = Q.T @ Yc
    cap = np.sum(proj * proj, axis=0) / np.sum(Yc * Yc, axis=0)
    cap[cap < max(CAP_THRESHOLD, 4.0 * (CAP_STATE_LAGS + 1) / X.shape[0])] = 0.0
    return cap


def capacity_emulation_analysis(seeds: Sequence[int]) -> pd.DataFrame:
    rows = []
    for seed in seeds:
        Up = probe_inputs(seed)
        Yc, deg = capacity_targets(Up, CAP_T0)
        d1 = np.flatnonzero(deg == 1)
        insts = {f: make_instance(f, seed, "A") for f in FAMILY_CODES}
        trajs = {f: simulate(insts[f], Up, range(N_OPS), "traj", ("probe", seed)) for f in FAMILY_CODES}
        prof = {f: np.stack([capacity_profile(trajs[f][:, :, op], Yc, CAP_T0) for op in range(N_OPS)]) for f in FAMILY_CODES}
        for s in FAMILY_CODES:
            for t in FAMILY_CODES:
                if s == t:
                    continue
                pad = rest_observable(insts[t])
                for op in range(N_OPS):
                    ps, pt = trajs[s][:, :, op], trajs[t][:, :, op]
                    mp = fit_lag_map(ps[:PROBE_TRAIN_EPISODES], pt[:PROBE_TRAIN_EPISODES], pad, FIR_TAPS, False)
                    ms, mt = prof[s][op][d1], prof[t][op][d1]
                    tot_s, tot_t = prof[s][op].sum(), prof[t][op].sum()
                    rows.append({"seed": seed, "source": s, "target": t, "op": op,
                                 "emulation_r2": r2_score_flat(ps[PROBE_TRAIN_EPISODES:], mp.apply(pt[PROBE_TRAIN_EPISODES:])),
                                 "target_total_capacity": tot_t, "source_total_capacity": tot_s,
                                 "memory_profile_corr": float(np.corrcoef(ms, mt)[0, 1]) if ms.std() > 0 and mt.std() > 0 else float("nan"),
                                 "lf_product": (ms.sum() / tot_s if tot_s > 0 else np.nan) * (mt.sum() / tot_t if tot_t > 0 else np.nan)})
    return pd.DataFrame(rows)


# -----------------------------------------------------------------------------
# Statistics
# -----------------------------------------------------------------------------


def hier_boot_fast(df: pd.DataFrame, col: str, seed: int) -> Tuple[float, float, float]:
    """Median of dataset medians; CI by resampling datasets, then rows within datasets."""
    groups = [g[col].dropna().to_numpy(np.float64) for _, g in df.groupby("dataset")]
    groups = [g for g in groups if len(g)]
    if not groups:
        return (float("nan"),) * 3
    obs = float(np.median([np.median(g) for g in groups]))
    rng = np.random.default_rng(seed)
    M = np.stack([np.median(g[rng.integers(0, len(g), (N_BOOT, len(g)))], axis=1) for g in groups])
    pick = rng.integers(0, len(groups), (N_BOOT, len(groups)))
    stats = np.median(M[pick, np.arange(N_BOOT)[:, None]], axis=1)
    return obs, float(np.quantile(stats, 0.025)), float(np.quantile(stats, 0.975))


def wilcoxon_two(x: np.ndarray, y: np.ndarray) -> Tuple[float, int]:
    d = np.asarray(x, float) - np.asarray(y, float)
    d = d[np.isfinite(d)]
    if len(d) < 3 or np.all(np.abs(d) < 1e-12):
        return float("nan"), len(d)
    return float(wilcoxon(d, zero_method="wilcox").pvalue), len(d)


def blocked_rhos(df: pd.DataFrame, x: str, y: str) -> np.ndarray:
    out = []
    for _, g in df.groupby("dataset"):
        g = g[[x, y]].dropna()
        if len(g) >= 5 and g[x].std() > 0 and g[y].std() > 0:
            out.append(float(spearmanr(g[x], g[y])[0]))
    return np.array(out)


def rho_test(rhos: np.ndarray, alternative: str) -> Dict[str, float]:
    if len(rhos) < 3:
        return {"mean_rho": float("nan"), "datasets": int(len(rhos)), "p": float("nan")}
    p = float(wilcoxon(rhos, alternative=alternative).pvalue) if np.any(rhos != 0) else 1.0
    return {"mean_rho": float(rhos.mean()), "median_rho": float(np.median(rhos)), "datasets": int(len(rhos)),
            "datasets_in_direction": int(np.sum(rhos < 0) if alternative == "less" else np.sum(rhos > 0)), "p": p}


def cluster_boot_spearman(df: pd.DataFrame, x: str, y: str, cluster: List[str], seed: int) -> Dict[str, float]:
    df = df.dropna(subset=[x, y])
    groups = [g for _, g in df.groupby(cluster)]
    obs = float(spearmanr(df[x], df[y])[0])
    rng = np.random.default_rng(seed)
    bs = []
    for _ in range(min(N_BOOT, 2000)):
        g = pd.concat([groups[i] for i in rng.integers(0, len(groups), len(groups))])
        bs.append(spearmanr(g[x], g[y])[0])
    bs = np.array(bs)
    return {"rho": obs, "ci95_low": float(np.nanquantile(bs, 0.025)), "ci95_high": float(np.nanquantile(bs, 0.975)),
            "p_one_sided_positive": float(np.mean(bs <= 0)), "n_points": int(len(df)), "n_clusters": len(groups)}


def lodo_prediction(df: pd.DataFrame, x: str, y: str) -> Tuple[pd.DataFrame, Dict[str, float]]:
    out = []
    for d in df.dataset.unique():
        tr = df[df.dataset != d].dropna(subset=[x, y])
        te = df[df.dataset == d].dropna(subset=[x, y])
        if len(tr) < 10 or len(te) == 0:
            continue
        coef = np.polyfit(tr[x], tr[y], 2)
        out.append(te.assign(predicted=np.polyval(coef, te[x])))
    if not out:
        return pd.DataFrame(), {}
    o = pd.concat(out)
    rh = blocked_rhos(o, "predicted", y)
    return o, {"pooled_spearman": float(spearmanr(o.predicted, o[y])[0]), "mae": float(np.mean(np.abs(o.predicted - o[y]))),
               "within_dataset": rho_test(rh, "greater")}





# =============================================================================
# v6: validation against measured devices
# =============================================================================

import zipfile
from scipy.optimize import least_squares

V5_DIR = PROJECT_ROOT / ("portability_v5_operating_point" + ("_ideal" if IDEAL else "") + ("_smoke" if SMOKE else ""))
MEAS_DIR_ENV = os.environ.get("MEMRISTIVE_MEASURED_DIR", "").strip()
TX_W = 100e-6            # pulse width stated in the Methods of Wang et al. 2026 (J. Phys. D 59 015101)
TX_INTERVALS = [8, 16, 32, 64, 128, 256, 512, 1024]
TX_AMPS = [1.5, 2.0, 2.5, 3.0, 3.5, 4.0]
TX_VREF = 4.0
TX_RESTARTS = 24 if not SMOKE else 4
TX_NC_LIST = [1, 2, 3]
TX_NC = 3
TX_LEVELS = [0.05, 0.50]
PAIR_METHODS = ["full_refit", "zero_shot", "moment", "paired_linear", "paired_rff", "distill_rff", "pseudo_label", "twin_paired_rff"]


def _ci_plot(ax, x, med, lo, hi, color, marker, label=None, ms=9):
    med, lo, hi = map(np.asarray, (med, lo, hi))
    ax.errorbar(x, med, yerr=[med - lo, hi - med], fmt=marker, color=color, ms=ms, capsize=5, lw=2, label=label)


def compare_ds(a: pd.Series, b: pd.Series, two_sided: bool = False, higher: bool = False, seed: int = 0) -> Dict[str, float]:
    j = pd.concat([a.rename("a"), b.rename("b")], axis=1).dropna()
    if two_sided:
        p = wilcoxon_two(j.a, j.b)[0]
    else:
        p = wilcoxon_less(j.b.to_numpy(), j.a.to_numpy())[0] if higher else wilcoxon_less(j.a.to_numpy(), j.b.to_numpy())[0]
    d = (j.a - j.b).to_numpy()
    rng = np.random.default_rng(seed)
    bs = np.median(d[rng.integers(0, len(d), (N_BOOT, len(d)))], axis=1) if len(d) else np.array([np.nan])
    better = (j.a > j.b) if higher else (j.a < j.b)
    return {"median_a": float(j.a.median()), "median_b": float(j.b.median()), "median_difference": float(np.median(d)) if len(d) else np.nan,
            "diff_ci95_low": float(np.quantile(bs, 0.025)), "diff_ci95_high": float(np.quantile(bs, 0.975)),
            "datasets_a_better": int(better.sum()), "datasets": int(len(j)), "p": p}


# -----------------------------------------------------------------------------
# Measured data
# -----------------------------------------------------------------------------


def find_measured_files() -> Tuple[Path, Optional[Path]]:
    roots = [Path(MEAS_DIR_ENV)] if MEAS_DIR_ENV else []
    roots += [PROJECT_ROOT / "data" / "measured", PROJECT_ROOT / "data", PROJECT_ROOT, Path("/content")]
    tiox_dir, xlsx = None, None
    for r in roots:
        if not r.exists():
            continue
        for depth_glob in ["*", "*/*", "*/*/*"]:
            for p in r.glob(depth_glob):
                if tiox_dir is None and p.name == "Dataset_Fig2.csv" and (p.parent / "Dataset_Fig3.csv").exists():
                    tiox_dir = p.parent
                if tiox_dir is None and p.name == "Dataset_Figures.zip":
                    out = DIRS["data"] / "tiox"
                    out.mkdir(parents=True, exist_ok=True)
                    with zipfile.ZipFile(p) as z:
                        for m in z.namelist():
                            if m.endswith(".csv") and "/" not in m.strip("/"):
                                (out / Path(m).name).write_bytes(z.read(m))
                    if (out / "Dataset_Fig2.csv").exists():
                        tiox_dir = out
                if xlsx is None and p.name == "Dataset_2.xlsx":
                    xlsx = p
            if tiox_dir is not None and xlsx is not None:
                break
        if tiox_dir is not None and xlsx is not None:
            break
    if tiox_dir is None:
        raise FileNotFoundError("TiOx files not found. Put Dataset_Figures.zip (or its CSVs) in "
                                f"{PROJECT_ROOT / 'data' / 'measured'} or set MEMRISTIVE_MEASURED_DIR.")
    return tiox_dir, xlsx


def _num(s) -> np.ndarray:
    return pd.to_numeric(s, errors="coerce").to_numpy(dtype=np.float64)


def load_tiox(d: Path) -> Dict[str, Any]:
    f2 = pd.read_csv(d / "Dataset_Fig2.csv", header=[0, 1], encoding="utf-8-sig")
    f3 = pd.read_csv(d / "Dataset_Fig3.csv", header=[0, 1], encoding="utf-8-sig")
    f4 = pd.read_csv(d / "Dataset_Fig4.csv", header=[0, 1], encoding="utf-8-sig")
    out: Dict[str, Any] = {"train_int": {}, "relax_int": {}, "train_amp": {}, "relax_amp": {}, "patterns": {}}
    for i in TX_INTERVALS:
        out["train_int"][i] = _num(f2[("Fig2_c", f"{i}us")])[:201]
    for a in TX_AMPS:
        out["train_amp"][a] = _num(f2[("Fig2_d", f"{a} V")])[:201]
    c3a = [c for c in f3.columns if c[0] == "Fig3_a"]
    for k, i in enumerate(TX_INTERVALS):
        t, y = _num(f3[c3a[2 * k]]), _num(f3[c3a[2 * k + 1]])
        m = np.isfinite(t) & np.isfinite(y)
        out["relax_int"][i] = (t[m], y[m])
    tb = _num(f3[("Fig3_b", "Time (s)")])
    for a in TX_AMPS:
        y = _num(f3[("Fig3_b", f"{a} V")])
        m = np.isfinite(tb) & np.isfinite(y)
        out["relax_amp"][a] = (tb[m], y[m])
    for cond in ["Fig4_c", "Fig4_d", "Fig4_e", "Fig4_f"]:
        cols = [c for c in f4.columns if c[0] == cond]
        M = np.stack([_num(f4[c]) for c in cols], axis=1)
        M = M[np.all(np.isfinite(M), axis=1)]
        out["patterns"][cond] = (M, [c[1] for c in cols])
    tt = _num(f2[("Fig2_b", "Time (s)")])
    tt = tt[np.isfinite(tt)]
    out["sampling_period_s"] = float(np.median(np.diff(tt[:200])))
    return out


def load_msio2(p: Optional[Path]) -> Dict[str, Any]:
    if p is None:
        return {}
    import openpyxl
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    res = {}

    def col(rows, i):
        return np.array([r[i] for r in rows[3:] if i < len(r) and isinstance(r[i], (int, float)) and r[i] > 0], float)
    if "Figure2" in wb.sheetnames:
        rows = list(wb["Figure2"].iter_rows(values_only=True))
        for blk, (h, l) in {"endurance_block_1": (24, 25), "endurance_block_2": (218, 219)}.items():
            H, Lr = col(rows, h), col(rows, l)
            if len(H) > 5 and len(Lr) > 5:
                res[blk] = {"cycles": int(min(len(H), len(Lr))), "HRS_sd_log10": float(np.log10(H).std(ddof=1)),
                            "LRS_sd_log10": float(np.log10(Lr).std(ddof=1)), "HRS_median_ohm": float(np.median(H)),
                            "LRS_median_ohm": float(np.median(Lr))}
    if "Figure4" in wb.sheetnames:
        rows = list(wb["Figure4"].iter_rows(values_only=True))
        for nm, i in [("HRS", 18), ("LRS_CC50uA", 19), ("LRS_CC100uA", 20), ("LRS_CC500uA", 21)]:
            v = col(rows, i)
            if len(v) > 5:
                res[f"compliance_{nm}"] = {"cycles": int(len(v)), "sd_log10": float(np.log10(v).std(ddof=1)), "median_ohm": float(np.median(v))}
    return res


# -----------------------------------------------------------------------------
# Calibration of a multi-timescale volatile model to the TiOx pulse data
# -----------------------------------------------------------------------------


def _rate(V: float, vth: float, alpha: float) -> float:
    return max(V - vth, 0.0) ** alpha


TX_DT_AMP = 8e-6            # interval of the amplitude series (Methods of Wang et al. 2026)
TX_READ_V = 0.1            # read voltage; also the "0" bit of the 3-bit inputs
TX_PATTERN_DELAYS = {"Fig4_c": 0.0, "Fig4_d": 108e-6, "Fig4_e": 648e-6, "Fig4_f": 1728e-6}


def tx_unpack(p: np.ndarray, nc: int) -> Dict[str, Any]:
    """Parameters of run A (interval series, 4 V, 100 us) plus the rate and output gain of run B (amplitude series)."""
    lkap, lm = p[:nc], p[nc:2 * nc]
    ltau0, dtau = p[2 * nc], p[2 * nc + 1: 3 * nc]
    vth, alpha, lcb, lsb = p[3 * nc: 3 * nc + 4]
    tau = 10 ** (ltau0 + np.concatenate([[0.0], np.cumsum(dtau)]))
    m = 10 ** lm
    k = 10 ** lkap * m / (TX_W * max(_rate(TX_VREF, vth, alpha), 1e-12))
    return {"k": k, "m": m, "tau": tau, "vth": vth, "alpha": alpha, "cb": 10 ** lcb, "sb": 10 ** lsb}


def tx_train(K: np.ndarray, M: np.ndarray, TAU: np.ndarray, dts: np.ndarray, N: int = 200) -> Tuple[np.ndarray, np.ndarray]:
    """Pulse trains: write for TX_W, then decay for the interval, read at the end of each interval."""
    y = np.zeros_like(K)
    out = np.zeros((K.shape[0], N + 1))
    e = np.exp(-K * TX_W / M)
    dd = np.exp(-dts[:, None] / TAU)
    for n in range(N):
        y = (M + (y - M) * e) * dd
        out[:, n + 1] = y.sum(1)
    return out, y


def tx_predict(q: Dict[str, Any], D: Dict[str, Any], keys: List[Tuple[str, Any]]) -> Dict[Tuple[str, Any], np.ndarray]:
    out = {}
    ki = [c for t, c in keys if t == "int"]
    ka = [c for t, c in keys if t == "amp"]
    if ki:
        K = np.tile(q["k"] * _rate(TX_VREF, q["vth"], q["alpha"]), (len(ki), 1))
        M = np.tile(q["m"], (len(ki), 1))
        tr, ye = tx_train(K, M, q["tau"][None, :], np.array(ki, float) * 1e-6)
        for r, i in enumerate(ki):
            I0 = D["train_int"][i][0]
            t = D["relax_int"][i][0]
            out[("int", i)] = I0 * (1 + tr[r])
            out[("int_relax", i)] = I0 * (1 + (ye[r][None, :] * np.exp(-t[:, None] / q["tau"][None, :])).sum(1))
    if ka:
        K = np.stack([q["k"] * q["cb"] * _rate(a, q["vth"], q["alpha"]) for a in ka])
        M = np.tile(q["m"], (len(ka), 1))
        tr, ye = tx_train(K, M, q["tau"][None, :], np.full(len(ka), TX_DT_AMP))
        for r, a in enumerate(ka):
            I0 = D["train_amp"][a][0]
            t = D["relax_amp"][a][0]
            out[("amp", a)] = I0 * (1 + q["sb"] * tr[r])
            out[("amp_relax", a)] = I0 * (1 + q["sb"] * (ye[r][None, :] * np.exp(-t[:, None] / q["tau"][None, :])).sum(1))
    return out


def tx_measured(D: Dict[str, Any], key: Tuple[str, Any]) -> np.ndarray:
    t, c = key
    if t == "int":
        return D["train_int"][c]
    if t == "amp":
        return D["train_amp"][c]
    if t == "int_relax":
        return D["relax_int"][c][1]
    return D["relax_amp"][c][1]


def tx_conditions() -> List[Tuple[str, Any]]:
    return [("int", i) for i in TX_INTERVALS] + [("amp", a) for a in TX_AMPS]


def tx_residual(p: np.ndarray, nc: int, D: Dict[str, Any], conds: List[Tuple[str, Any]]) -> np.ndarray:
    q = tx_unpack(p, nc)
    pr = tx_predict(q, D, conds)
    res = []
    for t, c in conds:
        for kk in [(t, c), (t + "_relax", c)]:
            y = tx_measured(D, kk)
            I0 = (D["train_int"][c] if t == "int" else D["train_amp"][c])[0]
            res.append((pr[kk] - y) / I0 * math.sqrt(201.0 / len(y)))
    return np.concatenate(res)


def tx_bounds(nc: int) -> Tuple[List[float], List[float]]:
    # Threshold between the 0.1 V read level and the lowest tested amplitude (1.5 V); 0.5 V lower bound.
    lb = [-4.0] * nc + [-3.0] * nc + [-6.0] + [0.3] * (nc - 1) + [0.5, 0.2, -3.0, -2.0]
    ub = [math.log10(5.0)] * nc + [0.5] * nc + [-3.0] + [4.0] * (nc - 1) + [1.45, 4.0, 3.0, 2.0]
    return lb, ub


def fit_tx(D: Dict[str, Any], nc: int, conds: Optional[List[Tuple[str, Any]]] = None, x0: Optional[np.ndarray] = None,
           restarts: int = TX_RESTARTS, seed: int = 0):
    conds = conds or tx_conditions()
    lb, ub = tx_bounds(nc)
    rng = np.random.default_rng(seed)
    starts = [x0] if x0 is not None else [np.array([rng.uniform(l, u) for l, u in zip(lb, ub)]) for _ in range(restarts)]
    best = None
    for s0 in starts:
        try:
            s = least_squares(tx_residual, np.clip(s0, lb, ub), bounds=(lb, ub), args=(nc, D, conds), max_nfev=3000, x_scale="jac")
        except Exception:
            continue
        if best is None or s.cost < best.cost:
            best = s
    return best


def r2(y: np.ndarray, p: np.ndarray) -> float:
    return float(1 - np.sum((y - p) ** 2) / max(np.sum((y - y.mean()) ** 2), 1e-30))


def tx_pattern_predict(params: Dict[str, float], delay: float) -> np.ndarray:
    """Response I/I(000) to the eight 3-bit inputs: '1' = 4 V and '0' = 0.1 V, 100 us pulses 8 us apart, read after `delay`."""
    m = np.array([params[f"m{j}"] for j in range(TX_NC)])
    k = np.array([params[f"k{j}"] for j in range(TX_NC)])
    tau = np.array([params[f"tau{j}"] for j in range(TX_NC)])
    out = []
    for i in range(8):
        y = np.zeros(TX_NC)
        for pos, bit in enumerate(f"{i:03b}"):
            v = TX_VREF if bit == "1" else TX_READ_V
            r = _rate(v, params["v_th"], params["alpha"])
            y = m + (y - m) * np.exp(-k * r * TX_W / m) if r > 0 else y * np.exp(-TX_W / tau)
            y = y * np.exp(-(8e-6 if pos < 2 else delay) / tau)
        out.append(1.0 + y.sum())
    out = np.array(out)
    return out / out[0]


def pattern_validation(D: Dict[str, Any], params: Dict[str, float]) -> Tuple[pd.DataFrame, Dict[str, Any]]:
    """Out-of-sample check: the 3-bit measurements were not used in the fit. One free amplitude scale is fitted,
    because the 3-bit runs show a much smaller response than the pulse-train runs."""
    rows = []
    for cond, delay in TX_PATTERN_DELAYS.items():
        M, labs = D["patterns"][cond]
        mu = M.mean(0)
        mu = mu / mu[0]
        se = M.std(0, ddof=1) / mu[0] / math.sqrt(M.shape[0])
        pr = tx_pattern_predict(params, delay)
        for i in range(8):
            rows.append({"condition": cond, "read_delay_us": delay * 1e6, "pattern": f"{i:03b}", "measured": float(mu[i]),
                         "measured_se": float(se[i]), "predicted_unscaled": float(pr[i])})
    df = pd.DataFrame(rows)
    a, b = df.measured - 1.0, df.predicted_unscaled - 1.0
    s = float(np.sum(a * b) / max(np.sum(b * b), 1e-30))
    df["predicted"] = 1.0 + s * b
    per = {c: float(spearmanr(g.measured, g.predicted)[0]) for c, g in df.groupby("condition") if g.measured.std() > 0}
    out = {"amplitude_scale": s, "pearson_r": float(np.corrcoef(a, b)[0, 1]), "r2_with_one_scale": r2(a.to_numpy(), s * b.to_numpy()),
           "spearman_pooled": float(spearmanr(a, b)[0]), "spearman_by_read_delay": per}
    return df, out


def calibrate_tx(D: Dict[str, Any]) -> Dict[str, Any]:
    conds = tx_conditions()
    sel = {}
    fits = {}
    for nc in TX_NC_LIST:
        log(f"Fitting the {nc}-timescale volatile model to the TiOx data ...")
        b = fit_tx(D, nc, seed=nc)
        fits[nc] = b
        n_pts = len(b.fun)
        rss = float(np.sum(b.fun ** 2))
        pr = tx_predict(tx_unpack(b.x, nc), D, conds)
        r2s = [r2(tx_measured(D, k), pr[k]) for k in pr]
        sel[nc] = {"parameters": int(len(b.x)), "rss": rss, "aic": float(n_pts * math.log(rss / n_pts) + 2 * len(b.x)),
                   "median_curve_r2": float(np.median(r2s)), "min_curve_r2": float(np.min(r2s))}
    b = fits[TX_NC]
    q = tx_unpack(b.x, TX_NC)
    pr = tx_predict(q, D, conds)
    curve = []
    for k in pr:
        curve.append({"series": k[0], "condition": k[1], "r2_in_sample": r2(tx_measured(D, k), pr[k])})
    log("Leave-one-condition-out validation ...")
    loco = {}
    for c in conds:
        rest = [x for x in conds if x != c]
        bc = fit_tx(D, TX_NC, rest, x0=b.x)
        pc = tx_predict(tx_unpack(bc.x, TX_NC), D, [c])
        for k in pc:
            loco[k] = r2(tx_measured(D, k), pc[k])
    for row in curve:
        row["r2_left_out"] = loco.get((row["series"], row["condition"]), np.nan)
    lb, ub = tx_bounds(TX_NC)
    at_bound = [i for i, (v, l, u) in enumerate(zip(b.x, lb, ub)) if abs(v - l) < 1e-3 or abs(v - u) < 1e-3]
    # The calibrated family is run A (interval series), whose protocol is fully documented.
    params = {"v_th": float(q["vth"]), "alpha": float(q["alpha"]), "x0": 0.0}
    for j in range(TX_NC):
        params[f"m{j}"] = float(q["m"][j])
        params[f"k{j}"] = float(q["k"][j])
        params[f"tau{j}"] = float(q["tau"][j])
    a8, b4 = D["train_int"][8], D["train_amp"][4.0]
    runs = {"run_A_I0_A": float(a8[0]), "run_B_I0_A": float(b4[0]), "baseline_ratio_A_over_B": float(a8[0] / b4[0]),
            "run_A_I200_over_I0": float(a8[-1] / a8[0]), "run_B_I200_over_I0": float(b4[-1] / b4[0]),
            "fitted_rate_ratio_B_over_A": float(q["cb"]), "fitted_output_gain_B_over_A": float(q["sb"]),
            "note": "Fig 2c at 8 us and Fig 2d at 4.0 V use the same nominal protocol (4 V, 100 us, 8 us)"}
    return {"q": q, "x": b.x, "params": params, "selection": sel, "curves": pd.DataFrame(curve), "pred": pr,
            "nuisance": runs, "parameters_at_bound": at_bound}


def observed_range(D: Dict[str, Any]) -> float:
    """Largest measured response (I/I0 - 1) in run A, the state range the calibrated device was seen to span."""
    return float(max(np.nanmax(y / y[0] - 1.0) for y in D["train_int"].values()))


def register_tx(params: Dict[str, float], xhi: float) -> None:
    FAMILIES["TX"] = Family("TX", "TiOx-calibrated volatile", "Calibrated to measurements, dataset 10.5258/SOTON/D3702",
                            "unipolar", TX_VREF, 0.9 * params["v_th"], False, True, 1e-7, 0.0, xhi, dict(params))
    T_SWITCH["TX"] = calibrate_switch_time("TX")
    T_RECOVER["TX"] = recovery_duration("TX")


# -----------------------------------------------------------------------------
# Noise relative to response range, measured and simulated
# -----------------------------------------------------------------------------


def measured_pattern_stats(D: Dict[str, Any]) -> Tuple[pd.DataFrame, pd.DataFrame]:
    rows, means = [], {}
    for cond, (M, labs) in D["patterns"].items():
        L = np.log10(M)
        sd = L.std(0, ddof=1)
        mu = L.mean(0)
        means[cond] = mu
        rng_ = mu.max() - mu.min()
        rows.append({"condition": cond, "repeats": int(L.shape[0]), "median_repeat_sd_decades": float(np.median(sd)),
                     "pattern_range_decades": float(rng_), "noise_to_range": float(np.median(sd) / rng_),
                     "min_adjacent_gap_over_sd": float(np.min(np.diff(np.sort(mu))) / np.median(sd))})
    keys = list(means)
    rk = [{"condition_a": a, "condition_b": b, "pattern_order_spearman": float(spearmanr(means[a], means[b])[0])}
          for i, a in enumerate(keys) for b in keys[i + 1:]]
    return pd.DataFrame(rows), pd.DataFrame(rk)


def simulated_pattern_stats(u: int, reps: int = 50) -> pd.DataFrame:
    pats = np.array([[2.0 * int(b) - 1.0 for b in f"{i:03b}"] for i in range(8)])
    U = np.repeat(pats, reps, axis=0)
    rows = []
    for f in FAMILY_CODES + ["TX"]:
        inst = nominal_instance(f)
        o = simulate(inst, U, [u], "traj", ("pattern_noise", f))[:, -1, 0].reshape(8, reps)
        sd = o.std(1, ddof=1)
        mu = o.mean(1)
        rows.append({"family": f, "median_repeat_sd_decades": float(np.median(sd)), "pattern_range_decades": float(mu.max() - mu.min()),
                     "noise_to_range": float(np.median(sd) / max(mu.max() - mu.min(), 1e-12))})
    return pd.DataFrame(rows)


# -----------------------------------------------------------------------------
# Portability experiments involving the calibrated family
# -----------------------------------------------------------------------------


def run_unit_v6(ds: "Dataset", seed: int, u: int, tx_noise: float) -> Tuple[pd.DataFrame, pd.DataFrame, Dict[str, float]]:
    C = ds.n_classes
    ytr, yte = ds.y_train, ds.y_test
    folds = make_folds(ytr, stable_seed("folds", ds.name, seed))
    ws = class_weights(ytr, C).sum()
    key_tr, key_te = ("task", ds.name, seed, "train"), ("task", ds.name, seed, "test")
    tx = make_instance("TX", seed, "A")
    Ftr = simulate(tx, ds.U_train, range(N_OPS), "features", key_tr)
    Fte = simulate(tx, ds.U_test, range(N_OPS), "features", key_te)
    cv = np.zeros(N_OPS); te = np.zeros(N_OPS); cache: Dict[Tuple, Tuple] = {}
    for op in range(N_OPS):
        k = op_key("TX", op)
        if k not in cache:
            cache[k] = select_and_test(Ftr[op], Fte[op], ytr, yte, C, folds)
        cv[op], _, te[op] = cache[k]
    land = pd.DataFrame({"dataset": ds.name, "tier": ds.tier, "seed": seed, "family": "TX", "op": range(N_OPS),
                         "gain": GAINS[OP_GAIN_IDX], "frac": FRACS[OP_FRAC_IDX], "cv_ba": cv, "test_ba": te,
                         "test_nba": [nba(v, C) for v in te]})
    del Ftr, Fte
    Fr_tr, Fr_te = features_from_traj(ds.U_train), features_from_traj(ds.U_test)
    _, _, raw_ba = select_and_test(Fr_tr, Fr_te, ytr, yte, C, folds)
    extra = {"dataset": ds.name, "seed": seed, "raw_nba": nba(raw_ba, C)}

    def pair_rows(src: Instance, tgt: Instance, direction: str, level: str, analysis: str, methods: List[str]) -> List[Dict[str, Any]]:
        Fs_tr = simulate(src, ds.U_train, [u], "features", key_tr)[0]
        Fs_te = simulate(src, ds.U_test, [u], "features", key_te)[0]
        Ft_tr = simulate(tgt, ds.U_train, [u], "features", key_tr)[0]
        Ft_te = simulate(tgt, ds.U_test, [u], "features", key_te)[0]
        _, ai_s, _ = select_and_test(Fs_tr, Fs_te, ytr, yte, C, folds)
        mu_s, sd_s = std_stats(Fs_tr)
        Zs_tr = (Fs_tr - mu_s) / sd_s
        ro_s = ridge_fit(Zs_tr, ytr, C, ALPHA_REL[ai_s] * ws)
        mu_t, sd_t = std_stats(Ft_tr)
        Zt_tr, Zt_te = (Ft_tr - mu_t) / sd_t, (Ft_te - mu_t) / sd_t
        _, _, full_ba = select_and_test(Ft_tr, Ft_te, ytr, yte, C, folds)
        pred = lambda Z: balanced_accuracy(yte, ro_s.predict(Z), C)
        sds = stable_seed("map6", ds.name, seed, src.iid, tgt.iid, analysis)
        res = {"full_refit": full_ba}
        if "zero_shot" in methods:
            res["zero_shot"] = pred((Ft_te - mu_s) / sd_s)
        res["moment"] = pred(Zt_te)
        if "paired_linear" in methods:
            res["paired_linear"] = pred(fit_feature_map(Zt_tr, Zs_tr, "lin", sds)(Zt_te))
        res["paired_rff"] = pred(fit_feature_map(Zt_tr, Zs_tr, "rff", sds + 1)(Zt_te))
        if "distill_rff" in methods:
            res["distill_rff"] = balanced_accuracy(yte, np.argmax(fit_feature_map(Zt_tr, Zs_tr @ ro_s.W + ro_s.b, "rff", sds + 2)(Zt_te), 1), C)
        res["pseudo_label"] = pseudo_label_ba(Zt_tr, Zt_te, ro_s.predict(Zs_tr), yte, C, sds)
        if "twin_paired_rff" in methods:
            Ftw = simulate(nominal_instance(src.family), ds.U_train, [u], "features", key_tr)[0]
            mtw, stw = std_stats(Ftw)
            res["twin_paired_rff"] = pred(fit_feature_map(Zt_tr, (Ftw - mtw) / stw, "rff", sds + 3)(Zt_te))
        cc = column_corr(Fs_tr, Ft_tr)
        fn = nba(full_ba, C)
        return [{"dataset": ds.name, "tier": ds.tier, "seed": seed, "source": src.family, "target": tgt.family, "direction": direction,
                 "level": level, "analysis": analysis, "op": u, "method": m, "ba": ba, "nba": nba(ba, C), "full_refit_nba": fn,
                 "gap": fn - nba(ba, C), "feature_corr_mean": float(np.nanmean(cc))} for m, ba in res.items()]

    pairs = [(tx, make_instance(t, seed, "A"), "tx_source", "cross") for t in FAMILY_CODES]
    pairs += [(make_instance(s, seed, "A"), tx, "tx_target", "cross") for s in FAMILY_CODES]
    pairs += [(tx, make_instance_sigma("TX", seed, "B" if abs(g - 0.05) < 1e-12 else f"B_sig{g:.2f}", g), "tx_within", f"d2d_{g:.2f}") for g in TX_LEVELS]
    rows = []
    for src, tgt, direction, level in pairs:
        rows += pair_rows(src, tgt, direction, level, "main", PAIR_METHODS)
        if seed == ROBUST_SEED:
            with noise_level(0.0, tx_noise):
                rows += pair_rows(src, tgt, direction, level, "measured_noise", ["full_refit", "moment", "paired_rff", "pseudo_label"])
    return pd.DataFrame(rows), land, extra


# -----------------------------------------------------------------------------
# Figures
# -----------------------------------------------------------------------------


def make_v6_figures(D: Dict[str, Any], cal: Dict[str, Any], R: Dict[str, Any]) -> None:
    pr = cal["pred"]
    fig, axes = plt.subplots(2, 2, figsize=(18, 13))
    cm = plt.get_cmap("viridis")
    specs = [("int", TX_INTERVALS, "Pulse number", "Current (µA)", lambda c: f"{c} µs"),
             ("int_relax", TX_INTERVALS, "Time after train (ms)", "Current (µA)", lambda c: f"{c} µs"),
             ("amp", TX_AMPS, "Pulse number", "Current (µA)", lambda c: f"{c:.1f} V"),
             ("amp_relax", TX_AMPS, "Time after train (ms)", "Current (µA)", lambda c: f"{c:.1f} V")]
    for ax, (ser, conds, xl, yl, lab) in zip(axes.ravel(), specs):
        allx, ally = [], []
        for i, c in enumerate(conds):
            col = cm(i / max(1, len(conds) - 1) * 0.9)
            y = tx_measured(D, (ser, c)) * 1e6
            if "relax" in ser:
                x = (D["relax_int"][c][0] if ser == "int_relax" else D["relax_amp"][c][0]) * 1e3
            else:
                x = np.arange(len(y))
            step = max(1, len(x) // 40)
            ax.plot(x[::step], y[::step], "o", ms=4, color=col, alpha=0.7)
            ax.plot(x, pr[(ser, c)] * 1e6, "-", lw=2, color=col, label=lab(c))
            allx += [x.min(), x.max()]; ally += [np.nanmin(y), np.nanmax(y)]
        set_bounds(ax, "x", min(allx), max(allx)); set_bounds(ax, "y", min(ally), max(ally))
        ax.set_xlabel(xl); ax.set_ylabel(yl)
        style_ax(ax); outlined_legend(ax, loc="best", fontsize=11, ncol=2)
    fig.tight_layout(h_pad=3, w_pad=3); panel_letters(fig, list(axes.ravel())); save_fig(fig, "fig15_tiox_calibration")

    fig, ax = plt.subplots(figsize=(8.6, 7.0))
    pv = R["pattern_validation"]
    for (cond, g), col in zip(pv.groupby("read_delay_us"), ["#1b6ca8", "#2a9d5c", "#e08a1e", "#c0392b"]):
        ax.errorbar(g.predicted, g.measured, yerr=1.96 * g.measured_se, fmt="o", ms=8, color=col, capsize=3, label=f"Read after {cond:.0f} µs")
    lo_ = min(pv.measured.min() - 1.96 * pv.measured_se.max(), pv.predicted.min())
    hi_ = max(pv.measured.max() + 1.96 * pv.measured_se.max(), pv.predicted.max())
    ax.plot([lo_, hi_], [lo_, hi_], "k--", lw=1.2)
    set_bounds(ax, "x", lo_, hi_); set_bounds(ax, "y", lo_, hi_)
    ax.set_xlabel("Predicted current, relative to input 000"); ax.set_ylabel("Measured current, relative to input 000")
    style_ax(ax); outlined_legend(ax, loc="upper left")
    fig.tight_layout(); save_fig(fig, "fig17_tiox_three_bit_prediction")

    fig, axes = plt.subplots(1, 2, figsize=(19, 6.6), gridspec_kw={"width_ratios": [1.0, 1.3]})
    ax = axes[0]
    ms_, sim = R["measured_patterns"], R["sim_patterns"]
    labels = [f"Condition {c[-1]}" for c in ms_.condition] + ["TiOx model" if f == "TX" else f for f in sim.family]
    vals = list(ms_.noise_to_range) + list(sim.noise_to_range)
    cols = ["#2a9d5c"] * len(ms_) + ["#1b6ca8"] * len(sim)
    ax.bar(range(len(vals)), vals, color=cols, edgecolor="black", lw=0.8)
    ax.set_xticks(range(len(vals))); ax.set_xticklabels(labels, rotation=45, ha="right", fontsize=13)
    ax.set_xlim(-0.6, len(vals) - 0.4)
    set_bounds(ax, "y", 0.0, max(vals))
    from matplotlib.patches import Patch
    ax.set_ylabel("Repeat noise / pattern range")
    style_ax(ax)
    outlined_legend(ax, loc="upper left", handles=[Patch(facecolor="#2a9d5c", edgecolor="black", label="Measured TiOx"),
                                                   Patch(facecolor="#1b6ca8", edgecolor="black", label="Simulated, nominal noise")])
    ax = axes[1]
    s = R["tx_summary"]
    s = s[s.analysis == "main"]
    order = [("tx_source", "cross", "TiOx as\nsource"), ("tx_target", "cross", "TiOx as\ntarget"),
             ("tx_within", "d2d_0.05", "TiOx 5%\nwithin"), ("tx_within", "d2d_0.50", "TiOx 50%\nwithin")]
    allv = [0.0]
    for (m, lab_, col, mk), off in zip([("moment", "Moment matching", "#e08a1e", "^"), ("paired_rff", "Paired nonlinear", "#1b6ca8", "o"),
                                         ("pseudo_label", "Pseudo labels", "#7b4fa0", "D"), ("twin_paired_rff", "Twin nonlinear", "#2a9d5c", "s")],
                                        [-0.18, -0.06, 0.06, 0.18]):
        g = [s[(s.direction == d) & (s.level == l) & (s.method == m)] for d, l, _ in order]
        med = [x.median_gap.iloc[0] if len(x) else np.nan for x in g]
        lo = [x.ci95_low.iloc[0] if len(x) else np.nan for x in g]
        hi = [x.ci95_high.iloc[0] if len(x) else np.nan for x in g]
        _ci_plot(ax, np.arange(len(order)) + off, med, lo, hi, col, mk, lab_, ms=8)
        allv += [v for v in lo + hi if np.isfinite(v)]
    ax.axhline(0, color="0.4", ls=":", lw=1)
    ax.set_xticks(range(len(order))); ax.set_xticklabels([o[2] for o in order], fontsize=13); ax.set_xlim(-0.5, len(order) - 0.5)
    set_bounds(ax, "y", min(allv), max(allv))
    ax.set_ylabel("Gap to full refit (normalized BA)")
    style_ax(ax); outlined_legend(ax, loc="upper right", fontsize=12)
    fig.tight_layout(w_pad=3); panel_letters(fig, axes); save_fig(fig, "fig16_measured_noise_and_tiox_transfer")




# =============================================================================
# v7: measured device pair (two nominally identical TiOx runs)
# =============================================================================

V6_DIR = PROJECT_ROOT / ("portability_v6_measured_validation" + ("_ideal" if IDEAL else "") + ("_smoke" if SMOKE else ""))
PAIR_SWEEP = [0.05, 0.20, 0.50, 1.00, 1.50] if not SMOKE else [0.05, 0.50]
V7_METHODS = ["full_refit", "zero_shot", "moment", "coral", "paired_linear", "paired_rff", "distill_rff", "pseudo_label",
              f"scratch_k{K_MAIN}", f"moment_warm_k{K_MAIN}"]
LADDER7 = [("zero_shot", "Zero\nshot"), ("moment", "Moment\nmatching"), ("coral", "CORAL"), ("paired_linear", "Paired\nlinear"),
           ("paired_rff", "Paired\nnonlinear"), ("distill_rff", "Score\ndistillation"), ("pseudo_label", "Pseudo\nlabels"),
           (f"scratch_k{K_MAIN}", f"Scratch,\n{K_MAIN} labels"), (f"moment_warm_k{K_MAIN}", f"Warm start,\n{K_MAIN} labels")]


def observed_range_b(D: Dict[str, Any]) -> float:
    return float(max(np.nanmax(y / y[0] - 1.0) for y in D["train_amp"].values()))


def register_measured_pair(params_a: Dict[str, float], D: Dict[str, Any], runs: Dict[str, float]) -> Dict[str, Any]:
    """TX = run A (interval series). TB = run B with its own interface. TS = run B driven with run A's protocol."""
    pa = dict(params_a)
    pa["log_g0"] = 0.0
    register_tx(pa, observed_range(D))
    FAMILIES["TX"].params["log_g0"] = 0.0
    cb, sb = runs["fitted_rate_ratio_B_over_A"], runs["fitted_output_gain_B_over_A"]
    pb = dict(pa)
    for j in range(TX_NC):
        pb[f"m{j}"] = pa[f"m{j}"] * sb
        pb[f"k{j}"] = pa[f"k{j}"] * sb * cb
    pb["log_g0"] = float(math.log10(runs["run_B_I0_A"] / runs["run_A_I0_A"]))
    xb = observed_range_b(D)
    for code, name in [("TB", "TiOx run B, own interface"), ("TS", "TiOx run B, run A protocol")]:
        FAMILIES[code] = Family(code, name, "Calibrated to measurements, dataset 10.5258/SOTON/D3702", "unipolar", TX_VREF,
                                0.9 * pb["v_th"], False, True, 1e-7, 0.0, xb, dict(pb))
    T_SWITCH["TB"] = calibrate_switch_time("TB")
    T_RECOVER["TB"] = recovery_duration("TB")
    T_SWITCH["TS"], T_RECOVER["TS"] = T_SWITCH["TX"], T_RECOVER["TX"]
    a, b = nominal_instance("TX"), nominal_instance("TB")
    keys = [k for k in pa if k not in ("x0", "log_g0")]
    lr = {k: math.log(abs(pb[k]) / abs(pa[k])) for k in keys if pa[k] != 0}
    return {"run_B_rate_ratio": cb, "run_B_output_gain": sb, "run_B_log10_baseline_offset": pb["log_g0"],
            "log_ratios": lr, "rms_log_ratio_all_parameters": float(np.sqrt(np.mean(np.square(list(lr.values()))))),
            "rms_log_ratio_changed_parameters": float(np.sqrt(np.mean(np.square([v for v in lr.values() if abs(v) > 1e-12])))),
            "T_switch": {"TX": T_SWITCH["TX"], "TB": T_SWITCH["TB"], "TS": T_SWITCH["TS"]},
            "observed_range": {"run_A": observed_range(D), "run_B": xb}}


def pair_methods(ds: "Dataset", seed: int, src: Instance, tgt: Instance, op_s: int, op_t: int, rng_k, analysis: str) -> Dict[str, Any]:
    C = ds.n_classes
    ytr, yte = ds.y_train, ds.y_test
    folds = make_folds(ytr, stable_seed("folds", ds.name, seed))
    ws = class_weights(ytr, C).sum()
    key_tr, key_te = ("task", ds.name, seed, "train"), ("task", ds.name, seed, "test")
    Fs_tr = simulate(src, ds.U_train, [op_s], "features", key_tr)[0]
    Fs_te = simulate(src, ds.U_test, [op_s], "features", key_te)[0]
    Ft_tr = simulate(tgt, ds.U_train, [op_t], "features", key_tr)[0]
    Ft_te = simulate(tgt, ds.U_test, [op_t], "features", key_te)[0]
    _, ai_s, src_ba = select_and_test(Fs_tr, Fs_te, ytr, yte, C, folds)
    mu_s, sd_s = std_stats(Fs_tr)
    Zs_tr = (Fs_tr - mu_s) / sd_s
    ro_s = ridge_fit(Zs_tr, ytr, C, ALPHA_REL[ai_s] * ws)
    mu_t, sd_t = std_stats(Ft_tr)
    Zt_tr, Zt_te = (Ft_tr - mu_t) / sd_t, (Ft_te - mu_t) / sd_t
    _, _, full_ba = select_and_test(Ft_tr, Ft_te, ytr, yte, C, folds)
    pred = lambda Z: balanced_accuracy(yte, ro_s.predict(Z), C)
    sds = stable_seed("map7", ds.name, seed, src.iid, tgt.iid, analysis, op_s)
    res = {"full_refit": full_ba, "zero_shot": pred((Ft_te - mu_s) / sd_s), "moment": pred(Zt_te),
           "coral": pred(Zt_te @ coral_matrix(Zs_tr, Zt_tr)),
           "paired_linear": pred(fit_feature_map(Zt_tr, Zs_tr, "lin", sds)(Zt_te)),
           "paired_rff": pred(fit_feature_map(Zt_tr, Zs_tr, "rff", sds + 1)(Zt_te)),
           "distill_rff": balanced_accuracy(yte, np.argmax(fit_feature_map(Zt_tr, Zs_tr @ ro_s.W + ro_s.b, "rff", sds + 2)(Zt_te), 1), C),
           "pseudo_label": pseudo_label_ba(Zt_tr, Zt_te, ro_s.predict(Zs_tr), yte, C, sds)}
    idx = np.concatenate([rng_k.permutation(np.flatnonzero(ytr == c))[:K_MAIN] for c in range(C)])
    res[f"scratch_k{K_MAIN}"] = balanced_accuracy(yte, ridge_fit(Zt_tr[idx], ytr[idx], C, LAMBDA_PRIOR).predict(Zt_te), C)
    res[f"moment_warm_k{K_MAIN}"] = balanced_accuracy(yte, ridge_fit(Zt_tr[idx], ytr[idx], C, LAMBDA_PRIOR, prior=ro_s.W).predict(Zt_te), C)
    cc = column_corr(Fs_tr, Ft_tr)
    ws_v, wt_v = ro_s.W.reshape(-1), ridge_fit(Zt_tr, ytr, C, ALPHA_REL[ai_s] * ws).W.reshape(-1)
    return {"res": res, "full_nba": nba(full_ba, C), "src_nba": nba(src_ba, C), "corr": float(np.nanmean(cc)),
            "flip": float(np.nanmean(cc < 0)), "cosine": float(ws_v @ wt_v / max(np.linalg.norm(ws_v) * np.linalg.norm(wt_v), 1e-20))}


def run_unit_v7(ds: "Dataset", seed: int, u: int, sel_op: Optional[int], tx_noise: float) -> pd.DataFrame:
    rows = []
    A, B, S = nominal_instance("TX"), nominal_instance("TB"), nominal_instance("TS")
    rng_k = np.random.default_rng(stable_seed("kshot7", ds.name, seed))
    ops = [("universal", u)] + ([("source_selected", sel_op)] if sel_op is not None and sel_op != u else [])
    combos = [("measured_pair_shared_protocol", A, S), ("measured_pair_shared_protocol", S, A),
              ("measured_pair_own_interface", A, B), ("measured_pair_own_interface", B, A)]
    combos += [(f"simulated_spread_{g:.2f}", A, make_instance_sigma("TX", seed, f"B_sig{g:.2f}", g)) for g in PAIR_SWEEP]
    for noise_name, (c2c, rd) in [("nominal_noise", (SIGMA_C2C, SIGMA_READ)), ("measured_noise", (0.0, tx_noise))]:
        for op_name, op in ops:
            if op_name != "universal" and noise_name != "nominal_noise":
                continue
            for pair_kind, src, tgt in combos:
                if pair_kind.startswith("simulated") and op_name != "universal":
                    continue
                with noise_level(c2c, rd):
                    r = pair_methods(ds, seed, src, tgt, op, op, rng_k, f"{noise_name}_{op_name}")
                for m, ba in r["res"].items():
                    rows.append({"dataset": ds.name, "tier": ds.tier, "seed": seed, "pair": pair_kind,
                                 "direction": f"{src.family}->{tgt.family}", "noise": noise_name, "op_policy": op_name, "op": op,
                                 "method": m, "ba": ba, "nba": nba(ba, ds.n_classes), "full_refit_nba": r["full_nba"],
                                 "gap": r["full_nba"] - nba(ba, ds.n_classes), "source_nba": r["src_nba"],
                                 "feature_corr_mean": r["corr"], "feature_sign_flip_fraction": r["flip"], "readout_cosine": r["cosine"]})
    return pd.DataFrame(rows)


def make_v7_figures(R: Dict[str, Any]) -> None:
    s = R["summary"]
    fig, axes = plt.subplots(1, 2, figsize=(21, 7.0), gridspec_kw={"width_ratios": [1.6, 1.0]})
    ax = axes[0]
    allv = [0.0]
    ms = [m for m, _ in LADDER7]
    for (pk, lab, col, mk), off in zip([("measured_pair_shared_protocol", "Same protocol", "#1b6ca8", "o"),
                                         ("measured_pair_own_interface", "Own interface", "#e08a1e", "s")], [-0.12, 0.12]):
        g = s[(s.pair == pk) & (s.noise == "nominal_noise") & (s.op_policy == "universal")].set_index("method").reindex(ms)
        _ci_plot(ax, np.arange(len(ms)) + off, g.median_gap, g.ci95_low, g.ci95_high, col, mk, lab)
        allv += list(g.ci95_low.dropna()) + list(g.ci95_high.dropna())
    ax.axhline(0, color="0.4", ls=":", lw=1)
    ax.set_xticks(range(len(ms))); ax.set_xticklabels([l for _, l in LADDER7], fontsize=13); ax.set_xlim(-0.5, len(ms) - 0.5)
    set_bounds(ax, "y", min(allv), max(allv))
    ax.set_ylabel("Gap to full refit (normalized BA)")
    style_ax(ax); outlined_legend(ax, loc="upper right")
    ax = axes[1]
    pm = R["pair_moment"]
    sim = pm[pm.pair.str.startswith("simulated")]
    cmap = plt.get_cmap("plasma")
    levels = sorted(sim.pair.unique())
    for i, lv in enumerate(levels):
        g = sim[sim.pair == lv]
        ax.scatter(g.feature_corr_mean, g.gap, s=22, alpha=0.6, color=cmap(0.05 + 0.75 * i / max(1, len(levels) - 1)),
                   label=f"Simulated, {float(lv.split('_')[-1]) * 100:.0f}%")
    for pk, col, mk, lab in [("measured_pair_shared_protocol", "#1b6ca8", "o", "Measured pair, same protocol"),
                             ("measured_pair_own_interface", "#e08a1e", "s", "Measured pair, own interface")]:
        g = pm[pm.pair == pk]
        ax.scatter(g.feature_corr_mean, g.gap, s=70, marker=mk, color=col, edgecolor="black", lw=0.7, label=lab)
    set_bounds(ax, "x", float(pm.feature_corr_mean.min()) - 0.02, 1.0)
    set_bounds(ax, "y", min(0.0, float(pm.gap.min())), float(pm.gap.max()))
    ax.set_xlabel("Mean feature correlation, source vs target"); ax.set_ylabel("Moment-matching gap")
    style_ax(ax); outlined_legend(ax, loc="upper right", fontsize=11)
    fig.tight_layout(w_pad=3); panel_letters(fig, axes); save_fig(fig, "fig18_measured_device_pair")


def main() -> None:
    t_start = time.time()
    log("=" * 100)
    log(f"MEMRISTIVE RESERVOIR PORTABILITY v7 MEASURED DEVICE PAIR | root={ROOT} | smoke={SMOKE}")
    log("Drive mounting is intentionally not performed by this script.")
    for fam in FAMILY_CODES:
        T_SWITCH[fam] = calibrate_switch_time(fam)
        T_RECOVER[fam] = recovery_duration(fam)
    tiox_dir, _ = find_measured_files()
    D = load_tiox(tiox_dir)
    cal_path = V6_DIR / "config" / "tiox_calibration.json"
    if cal_path.exists():
        saved = json.loads(cal_path.read_text())
        params, runs = saved["params"], saved["nuisance"]
        log(f"Using the v6 calibration from {cal_path}")
    else:
        log("v6 calibration not found; calibrating now.")
        cal = calibrate_tx(D)
        params, runs = cal["params"], cal["nuisance"]
    if "fitted_rate_ratio_B_over_A" not in runs:
        raise RuntimeError("The saved calibration predates the documented-protocol fit. Delete it and rerun v6 or this script.")
    info = register_measured_pair(params, D, runs)
    save_json(info, DIRS["config"] / "measured_pair.json")
    log(f"Measured pair: run B rate x{info['run_B_rate_ratio']:.3f}, output gain x{info['run_B_output_gain']:.3f}, "
        f"baseline offset {info['run_B_log10_baseline_offset']:.3f} decades; T_switch {info['T_switch']}")
    ujson = PROJECT_ROOT / ("portability_v5_operating_point" + ("_smoke" if SMOKE else "")) / "config" / "universal_operating_point.json"
    u = int(json.loads(ujson.read_text())["op"]) if ujson.exists() else int(N_OPS - 3)
    sel = {}
    tl = V6_DIR / "tables" / "tiox_landscape.csv"
    if tl.exists():
        L = pd.read_csv(tl)
        for (d, s), g in L.groupby(["dataset", "seed"]):
            sel[(d, int(s))] = int(g.sort_values("op").cv_ba.to_numpy().argmax())
    mpth = V6_DIR / "tables" / "measured_pattern_noise.csv"
    tx_noise = float(pd.read_csv(mpth).median_repeat_sd_decades.median()) if mpth.exists() else float(measured_pattern_stats(D)[0].median_repeat_sd_decades.median())

    names = (["SmokeA", "SmokeB", "SmokeC"] if SMOKE else PRIOR_DATASETS + FRESH_DATASETS + HELD_OUT_DATASETS)
    if DATASET_FILTER:
        names = [d for d in names if d in DATASET_FILTER]
    loaded = []
    for name in names:
        try:
            ds = synthetic_dataset(name, stable_seed(name)) if SMOKE else load_dataset(name)
        except Exception as exc:
            log(f"SKIP {name}: {type(exc).__name__}: {exc}")
            continue
        if name in HELD_OUT_DATASETS or (SMOKE and name == "SmokeC"):
            ds.tier = "held_out"
        loaded.append(ds)
    loaded.sort(key=lambda d: (len(d.y_train) + len(d.y_test)) * d.U_train.shape[1])
    h = hashlib.sha256(json.dumps([params, runs, PAIR_SWEEP, u], sort_keys=True, default=str).encode()).hexdigest()[:10]
    U = []
    for ds in loaded:
        for seed in V4_SEEDS:
            pth = DIRS["units"] / f"v7_{ds.name}_s{seed}_{h}.csv"
            if pth.exists():
                U.append(pd.read_csv(pth)); continue
            t0 = time.time()
            try:
                df = run_unit_v7(ds, seed, u, sel.get((ds.name, seed)), tx_noise)
            except Exception:
                log(f"FAILED {ds.name} seed {seed}:\n{traceback.format_exc()}")
                continue
            df.to_csv(pth, index=False); U.append(df)
            g = df[(df.pair == "measured_pair_shared_protocol") & (df.noise == "nominal_noise") & (df.op_policy == "universal")].groupby("method").gap.median()
            log(f"{ds.name} seed {seed}: {(time.time() - t0) / 60:.1f} min | measured pair gap moment={g.get('moment', np.nan):.3f} "
                f"paired={g.get('paired_rff', np.nan):.3f} pseudo={g.get('pseudo_label', np.nan):.3f} zero_shot={g.get('zero_shot', np.nan):.3f}")
    units = pd.concat(U, ignore_index=True)
    units.to_csv(DIRS["tables"] / "measured_pair_units.csv", index=False)

    keys = ["dataset", "tier", "pair", "noise", "op_policy", "method"]
    cols = ["gap", "nba", "full_refit_nba", "source_nba", "feature_corr_mean", "feature_sign_flip_fraction", "readout_cosine"]
    # Pair = both directions pooled by seed and direction medians.
    pair = units.groupby(keys + ["direction"], as_index=False)[cols].median().groupby(keys, as_index=False)[cols].median()
    pair.to_csv(DIRS["tables"] / "measured_pair_dataset_level.csv", index=False)
    srows = []
    for k, g in pair.groupby(["pair", "noise", "op_policy", "method"]):
        med, lo, hi = hier_boot_fast(g, "gap", stable_seed("v7", *k))
        srows.append({"pair": k[0], "noise": k[1], "op_policy": k[2], "method": k[3], "median_gap": med, "ci95_low": lo, "ci95_high": hi,
                      "median_nba": float(g.nba.median()), "median_full_refit_nba": float(g.full_refit_nba.median())})
    summary = pd.DataFrame(srows)
    summary.to_csv(DIRS["tables"] / "measured_pair_summary.csv", index=False)
    diag = pair[pair.method == "moment"].groupby(["pair", "noise", "op_policy"])[["feature_corr_mean", "feature_sign_flip_fraction", "readout_cosine",
                                                                                 "full_refit_nba", "source_nba"]].median()
    sub = lambda pk, noise="nominal_noise", op="universal", m="moment": pair[(pair.pair == pk) & (pair.noise == noise) & (pair.op_policy == op)
                                                                              & (pair.method == m)].set_index("dataset").gap
    tests = {}
    for pk, tag in [("measured_pair_shared_protocol", "same protocol"), ("measured_pair_own_interface", "own interface")]:
        tests[f"M1 {tag}: paired nonlinear < moment"] = compare_ds(sub(pk, m="paired_rff"), sub(pk), seed=1)
        tests[f"M2 {tag}: pseudo labels < moment"] = compare_ds(sub(pk, m="pseudo_label"), sub(pk), seed=2)
        tests[f"M3 {tag}: moment < zero shot"] = compare_ds(sub(pk), sub(pk, m="zero_shot"), seed=3)
        tests[f"M4 {tag}: paired nonlinear vs pseudo labels (two-sided)"] = compare_ds(sub(pk, m="paired_rff"), sub(pk, m="pseudo_label"), True, seed=4)
    ph = holm({k: v["p"] for k, v in tests.items()})
    for k in tests:
        tests[k]["p_holm"] = ph[k]
    pm = pair[(pair.method == "moment") & (pair.noise == "nominal_noise") & (pair.op_policy == "universal")]
    simm = pm[pm.pair.str.startswith("simulated")]
    pos = {}
    if len(simm) >= 5:
        coef = np.polyfit(simm.feature_corr_mean, simm.gap, 2)
        for pk in ["measured_pair_shared_protocol", "measured_pair_own_interface"]:
            g = pm[pm.pair == pk]
            pos[pk] = {"median_feature_corr": float(g.feature_corr_mean.median()), "observed_median_gap": float(g.gap.median()),
                       "gap_predicted_from_simulated_spread_curve": float(np.median(np.polyval(coef, g.feature_corr_mean))),
                       "nearest_simulated_level_by_feature_corr": min(((abs(float(simm[simm.pair == lv].feature_corr_mean.median()) - float(g.feature_corr_mean.median())), lv)
                                                                       for lv in simm.pair.unique()))[1]}
    v5p = PROJECT_ROOT / ("portability_v5_operating_point" + ("_smoke" if SMOKE else "")) / "tables" / "all_units_universal_point.csv"
    if v5p.exists():
        v5 = pd.read_csv(v5p)
        v5m = v5[v5.method == "moment"].dropna(subset=["feature_corr_mean", "gap"])
        coef5 = np.polyfit(v5m.feature_corr_mean, v5m.gap, 2)
        for pk in pos:
            g = pm[pm.pair == pk]
            pos[pk]["gap_predicted_from_v5_curve"] = float(np.median(np.polyval(coef5, g.feature_corr_mean)))
    sel_cmp = {}
    for pk in ["measured_pair_shared_protocol", "measured_pair_own_interface"]:
        a = pair[(pair.pair == pk) & (pair.noise == "nominal_noise") & (pair.op_policy == "universal") & (pair.method == "full_refit")].set_index("dataset").nba
        b = pair[(pair.pair == pk) & (pair.noise == "nominal_noise") & (pair.op_policy == "source_selected") & (pair.method == "full_refit")].set_index("dataset").nba
        if len(b):
            sel_cmp[pk] = compare_ds(a, b, True, True, seed=9)
    R = {"summary": summary, "pair_moment": pm}
    try:
        make_v7_figures(R)
    except Exception:
        log("Figure generation failed:\n" + traceback.format_exc())

    elapsed = (time.time() - t_start) / 60
    pd.set_option("display.width", 240); pd.set_option("display.max_columns", 40); pd.set_option("display.max_rows", 400)
    pd.set_option("display.float_format", lambda v: f"{v:.4f}")
    J = lambda o: json.dumps(o, indent=1, default=str)
    print("\n" + "=" * 110)
    print("COPY-PASTE REPORT START - MEMRISTIVE RESERVOIR PORTABILITY v7 MEASURED DEVICE PAIR")
    print("=" * 110)
    print(f"Runtime: {elapsed:.1f} min | datasets {len(loaded)} | seeds {V4_SEEDS} | universal op {u} | measured noise {tx_noise:.5f} decades")
    print("\nMEASURED PAIR (run A = Fig 2c interval series, run B = Fig 2d amplitude series, same nominal protocol):", J(info))
    print("\nDIAGNOSTICS (moment rows, medians)"); print(diag.to_string())
    print("\nSUMMARY (gap to full refit, pooled over both directions)")
    print(summary.sort_values(["pair", "noise", "op_policy", "median_gap"]).to_string(index=False))
    print("\nTESTS (dataset level, Holm over all M tests)")
    print(pd.DataFrame([{"test": k, **v} for k, v in tests.items()]).to_string(index=False))
    print("\nWHERE THE MEASURED PAIR SITS ON THE SIMULATED SPREAD AXIS:", J(pos))
    print("\nFULL REFIT AT THE UNIVERSAL VS THE SOURCE-SELECTED POINT:", J(sel_cmp))
    by_tier = {t: compare_ds(g[g.method == "paired_rff"].set_index("dataset").gap, g[g.method == "moment"].set_index("dataset").gap)
               for t, g in pair[(pair.pair == "measured_pair_shared_protocol") & (pair.noise == "nominal_noise") & (pair.op_policy == "universal")].groupby("tier")}
    print("\nBY TIER, same protocol, paired nonlinear vs moment:", J(by_tier))
    print("=" * 110)
    print("COPY-PASTE REPORT END")
    print("=" * 110)
    log(f"Done in {elapsed:.1f} min. Tables: {DIRS['tables']} | Figures: {DIRS['figures']}")


if __name__ == "__main__":
    main()
